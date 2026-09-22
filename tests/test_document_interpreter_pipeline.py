import io
import sys
import unittest
from pathlib import Path


BACKEND_DIR = Path(__file__).resolve().parents[1]
if str(BACKEND_DIR) not in sys.path:
    sys.path.insert(0, str(BACKEND_DIR))

from document_extractors import DocumentExtractionError
from document_hierarchy_contracts import CONFIDENCE_HIGH
from document_interpreter_pipeline import interpret_document_to_chunks
from document_interpreter_contract import INTERPRETER_STATUS_ACCEPTED


def _require_docx():
    try:
        from docx import Document
    except ImportError as exc:
        raise unittest.SkipTest("python-docx is not installed") from exc
    return Document


def _require_pptx():
    try:
        from pptx import Presentation
    except ImportError as exc:
        raise unittest.SkipTest("python-pptx is not installed") from exc
    return Presentation


def _require_pypdf():
    try:
        from pypdf import PdfWriter
        from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject
    except ImportError as exc:
        raise unittest.SkipTest("pypdf is not installed") from exc
    return PdfWriter, DecodedStreamObject, DictionaryObject, NameObject


def _docx_bytes(document) -> bytes:
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _literal(text: str) -> str:
    return text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


def _text_command(text: str, *, x: int, y: int, size: int = 12, font: str = "F1") -> str:
    return f"BT /{font} {size} Tf {x} {y} Td ({_literal(text)}) Tj ET"


def _pdf_bytes(pages) -> bytes:
    PdfWriter, DecodedStreamObject, DictionaryObject, NameObject = _require_pypdf()
    writer = PdfWriter()
    for commands in pages:
        page = writer.add_blank_page(width=612, height=792)
        fonts = DictionaryObject({
            NameObject("/F1"): DictionaryObject({
                NameObject("/Type"): NameObject("/Font"),
                NameObject("/Subtype"): NameObject("/Type1"),
                NameObject("/BaseFont"): NameObject("/Helvetica"),
            }),
            NameObject("/F2"): DictionaryObject({
                NameObject("/Type"): NameObject("/Font"),
                NameObject("/Subtype"): NameObject("/Type1"),
                NameObject("/BaseFont"): NameObject("/Helvetica-Bold"),
            }),
        })
        page[NameObject("/Resources")] = DictionaryObject({
            NameObject("/Font"): fonts,
        })
        stream = DecodedStreamObject()
        stream.set_data("\n".join(commands).encode("utf-8"))
        page[NameObject("/Contents")] = writer._add_object(stream)
    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def _pptx_bytes(presentation) -> bytes:
    buffer = io.BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


class DocumentInterpreterPipelineTests(unittest.TestCase):
    def test_docx_file_runs_from_raw_bytes_to_section_owned_chunks(self):
        Document = _require_docx()
        document = Document()
        document.add_heading("1 Cellular Biology", level=1)
        document.add_paragraph("Cells are the basic unit of life. " * 20)
        document.add_heading("1.1 Membranes", level=2)
        document.add_paragraph("Membranes regulate transport and signaling. " * 20)

        result = interpret_document_to_chunks(
            _docx_bytes(document),
            "biology.docx",
            document_id="doc-1",
        )

        self.assertEqual(result.interpretation.status, INTERPRETER_STATUS_ACCEPTED)
        self.assertEqual(result.interpretation.confidence, CONFIDENCE_HIGH)
        self.assertEqual(
            [section.title for section in result.interpretation.sections],
            ["1 Cellular Biology", "1.1 Membranes"],
        )
        self.assertEqual(
            [chunk.section_path for chunk in result.chunks],
            ["1 Cellular Biology", "1 Cellular Biology > 1.1 Membranes"],
        )
        self.assertIn("Cells are the basic unit", result.chunks[0].chunk_text)
        self.assertIn("Membranes regulate", result.chunks[1].chunk_text)

    def test_pdf_file_runs_from_raw_bytes_to_section_owned_chunks(self):
        pdf = _pdf_bytes([[
            _text_command("1 Cardiac physiology", x=72, y=720, size=16, font="F2"),
            _text_command("Cardiac output depends on heart rate and stroke volume. " * 20, x=72, y=690, size=12),
        ]])

        result = interpret_document_to_chunks(
            pdf,
            "physiology.pdf",
            document_id="doc-1",
        )

        self.assertEqual(result.interpretation.status, INTERPRETER_STATUS_ACCEPTED)
        self.assertEqual([section.title for section in result.interpretation.sections], ["1 Cardiac physiology"])
        self.assertGreaterEqual(len(result.chunks), 1)
        self.assertEqual({chunk.section_id for chunk in result.chunks}, {"pdf-section-0001"})
        self.assertEqual({chunk.section_path for chunk in result.chunks}, {"1 Cardiac physiology"})
        self.assertIn(
            "Cardiac output depends",
            " ".join(chunk.chunk_text for chunk in result.chunks),
        )

    def test_pptx_file_runs_from_raw_bytes_to_slide_owned_chunks(self):
        Presentation = _require_pptx()
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = "Neural signaling"
        slide.placeholders[1].text = "Action potentials transmit information. " * 20

        result = interpret_document_to_chunks(
            _pptx_bytes(presentation),
            "slides.pptx",
            document_id="doc-1",
        )

        self.assertEqual(result.interpretation.status, INTERPRETER_STATUS_ACCEPTED)
        self.assertEqual([section.title for section in result.interpretation.sections], ["Neural signaling"])
        self.assertEqual(len(result.chunks), 1)
        self.assertEqual(result.chunks[0].section_id, "pptx-slide-0001")
        self.assertEqual(result.chunks[0].section_path, "Neural signaling")
        self.assertIn("Action potentials transmit", result.chunks[0].chunk_text)

    def test_unknown_format_fails_clearly(self):
        with self.assertRaisesRegex(DocumentExtractionError, "Unsupported document format"):
            interpret_document_to_chunks(
                b"plain text",
                "notes.txt",
                document_id="doc-1",
            )


if __name__ == "__main__":
    unittest.main()
