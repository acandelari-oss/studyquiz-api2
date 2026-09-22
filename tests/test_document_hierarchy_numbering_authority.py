import io
import sys
import unittest
from pathlib import Path
from typing import Optional


BACKEND_DIR = Path(__file__).resolve().parents[1]
if str(BACKEND_DIR) not in sys.path:
    sys.path.insert(0, str(BACKEND_DIR))

from canonical_document import SourcePosition, adapt_docx_document
from document_hierarchy_consistency import analyze_docx_heading_numbering_consistency
from document_hierarchy_contracts import (
    CONFIDENCE_HIGH,
    CONTENT_KIND_BODY,
    CONTENT_KIND_DOCUMENT_BOUNDARY,
    CONTENT_KIND_LIST_ITEM,
    CONTENT_KIND_TABLE,
    EVIDENCE_ROLE_BODY,
    EVIDENCE_ROLE_HEADING,
    EVIDENCE_ROLE_LIST,
    EVIDENCE_ROLE_TABLE,
    SOURCE_FORMAT_DOCX,
    SourceSpan,
    StructuralEvidence,
)
from document_hierarchy_evidence_mapper import map_document_to_structural_evidence
from document_hierarchy_explicit_preserver import preserve_explicit_docx_hierarchy
from document_hierarchy_numbering_authority import (
    AUTHORITY_AMBIGUOUS_NUMBERING,
    AUTHORITY_EXPLICIT_CONFLICT,
    AUTHORITY_EXPLICIT_CONFIRMED,
    AUTHORITY_LOCAL_OR_UNSUPPORTED,
    AUTHORITY_TRUSTED_NUMBERING_SYSTEM,
    classify_docx_numbering_authority,
    render_docx_numbering_authority,
)
from document_hierarchy_numbering_observer import observe_docx_numbering_patterns


def span(order: int) -> SourceSpan:
    position = SourcePosition("document_flow", 1, order, 0, 10)
    return SourceSpan(start=position, end=position)


def heading(
    order: int,
    *,
    level: int,
    ilvl: Optional[int] = None,
    num_id=5,
) -> StructuralEvidence:
    numbering = {}
    if ilvl is not None:
        numbering = {
            "source": "direct",
            "num_id": num_id,
            "abstract_num_id": 7,
            "ilvl": ilvl,
            "list_level": ilvl,
        }
    return StructuralEvidence(
        evidence_id=f"ev-{order:04d}",
        source_format=SOURCE_FORMAT_DOCX,
        source_order=order,
        source_span=span(order),
        raw_text=f"Heading {order}",
        normalized_text=f"Heading {order}",
        evidence_role=EVIDENCE_ROLE_HEADING,
        content_kind=CONTENT_KIND_DOCUMENT_BOUNDARY,
        native_evidence={
            "role_hint": "heading",
            "style_hint": f"Heading {level}",
            "native_hierarchy_hint": level,
        },
        numbering_evidence=numbering,
        confidence_hint=CONFIDENCE_HIGH,
    )


def numbered(
    order: int,
    *,
    ilvl: int,
    num_id=5,
    evidence_role=EVIDENCE_ROLE_LIST,
    content_kind=CONTENT_KIND_LIST_ITEM,
) -> StructuralEvidence:
    return StructuralEvidence(
        evidence_id=f"ev-{order:04d}",
        source_format=SOURCE_FORMAT_DOCX,
        source_order=order,
        source_span=span(order),
        raw_text=f"Numbered {order}",
        normalized_text=f"Numbered {order}",
        evidence_role=evidence_role,
        content_kind=content_kind,
        native_evidence={"role_hint": "list", "native_hierarchy_hint": ilvl},
        numbering_evidence={
            "source": "direct",
            "num_id": num_id,
            "abstract_num_id": 7,
            "ilvl": ilvl,
            "list_level": ilvl,
            "format": "decimal",
        },
        confidence_hint=CONFIDENCE_HIGH,
    )


def body(order: int) -> StructuralEvidence:
    return StructuralEvidence(
        evidence_id=f"ev-{order:04d}",
        source_format=SOURCE_FORMAT_DOCX,
        source_order=order,
        source_span=span(order),
        raw_text=f"Body {order}",
        normalized_text=f"Body {order}",
        evidence_role=EVIDENCE_ROLE_BODY,
        content_kind=CONTENT_KIND_BODY,
        native_evidence={"role_hint": "body"},
    )


def classify(items):
    structure = preserve_explicit_docx_hierarchy(items, document_id="doc")
    consistency = analyze_docx_heading_numbering_consistency(items, structure)
    observations = observe_docx_numbering_patterns(items)
    authority = classify_docx_numbering_authority(
        items,
        structure,
        consistency,
        observations,
    )
    return structure, consistency, observations, authority


def _require_docx():
    try:
        from docx import Document
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
    except ImportError as exc:  # pragma: no cover - dependency check.
        raise unittest.SkipTest("python-docx is not installed") from exc
    return Document, OxmlElement, qn


def _set_direct_numbering(paragraph, *, num_id: int = 5, ilvl: int = 0) -> None:
    _Document, OxmlElement, qn = _require_docx()
    p_pr = paragraph._p.get_or_add_pPr()  # noqa: SLF001 - test fixture XML.
    num_pr = p_pr.numPr
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        p_pr.append(num_pr)
    ilvl_el = num_pr.find(qn("w:ilvl"))
    if ilvl_el is None:
        ilvl_el = OxmlElement("w:ilvl")
        num_pr.append(ilvl_el)
    ilvl_el.set(qn("w:val"), str(ilvl))
    num_id_el = num_pr.find(qn("w:numId"))
    if num_id_el is None:
        num_id_el = OxmlElement("w:numId")
        num_pr.append(num_id_el)
    num_id_el.set(qn("w:val"), str(num_id))


class DocxNumberingAuthorityTests(unittest.TestCase):
    def test_explicit_heading_1_matching_numbering_is_confirmed(self):
        _structure, _consistency, _observations, authority = classify([
            heading(1, level=1, ilvl=0),
        ])

        self.assertEqual(authority.records[0].authority_class, AUTHORITY_EXPLICIT_CONFIRMED)
        self.assertEqual(authority.records[0].basis["heading_level"], 1)
        self.assertEqual(authority.records[0].basis["numbering_level"], 1)

    def test_explicit_heading_2_matching_numbering_is_confirmed(self):
        _structure, _consistency, _observations, authority = classify([
            heading(1, level=1, ilvl=0),
            heading(2, level=2, ilvl=1),
        ])

        self.assertEqual(
            [record.authority_class for record in authority.records],
            [AUTHORITY_EXPLICIT_CONFIRMED, AUTHORITY_EXPLICIT_CONFIRMED],
        )
        self.assertEqual(authority.records[1].basis["heading_level"], 2)
        self.assertEqual(authority.records[1].basis["numbering_level"], 2)

    def test_explicit_heading_conflict_is_diagnostic_only(self):
        structure, _consistency, _observations, authority = classify([
            heading(1, level=1, ilvl=2),
        ])

        self.assertEqual(authority.records[0].authority_class, AUTHORITY_EXPLICIT_CONFLICT)
        self.assertEqual(structure.nodes[0].level, 1)
        self.assertIsNone(structure.nodes[0].parent_id)

    def test_explicit_unnumbered_heading_does_not_create_false_conflict(self):
        _structure, _consistency, _observations, authority = classify([
            heading(1, level=1),
        ])

        self.assertEqual(authority.records, [])
        self.assertEqual(authority.diagnostics["explicit_conflict_count"], 0)

    def test_all_heading_1_multi_level_numbering_reports_confirmed_and_conflicts_only(self):
        structure, _consistency, _observations, authority = classify([
            heading(1, level=1, ilvl=0),
            heading(2, level=1, ilvl=1),
            heading(3, level=1, ilvl=2),
        ])

        self.assertEqual([node.level for node in structure.nodes], [1, 1, 1])
        self.assertEqual([node.parent_id for node in structure.nodes], [None, None, None])
        self.assertEqual(
            [record.authority_class for record in authority.records],
            [
                AUTHORITY_EXPLICIT_CONFIRMED,
                AUTHORITY_EXPLICIT_CONFLICT,
                AUTHORITY_EXPLICIT_CONFLICT,
            ],
        )

    def test_coherent_normal_numbering_is_ambiguous_and_creates_no_nodes(self):
        structure, _consistency, observations, authority = classify([
            numbered(1, ilvl=0),
            numbered(2, ilvl=1),
            numbered(3, ilvl=1),
            numbered(4, ilvl=0),
        ])

        self.assertEqual(structure.nodes, [])
        self.assertEqual(len(observations.observations), 4)
        self.assertEqual(
            {record.authority_class for record in authority.records},
            {AUTHORITY_AMBIGUOUS_NUMBERING},
        )
        self.assertEqual(authority.diagnostics["trusted_numbering_system_count"], 0)

    def test_separate_native_numbering_groups_remain_traceable(self):
        _structure, _consistency, _observations, authority = classify([
            numbered(1, ilvl=0, num_id=5),
            numbered(2, ilvl=0, num_id=6),
        ])

        self.assertEqual([record.numbering_group_id for record in authority.records], ["num:5", "num:6"])

    def test_local_list_evidence_stays_non_hierarchical(self):
        structure, _consistency, _observations, authority = classify([
            body(1),
            numbered(2, ilvl=0),
            numbered(3, ilvl=0),
        ])

        self.assertEqual(structure.nodes, [])
        self.assertEqual(
            [record.authority_class for record in authority.records],
            [AUTHORITY_AMBIGUOUS_NUMBERING, AUTHORITY_AMBIGUOUS_NUMBERING],
        )

    def test_table_contained_numbering_is_local_or_unsupported_when_preserved(self):
        structure, _consistency, _observations, authority = classify([
            numbered(
                1,
                ilvl=0,
                evidence_role=EVIDENCE_ROLE_TABLE,
                content_kind=CONTENT_KIND_TABLE,
            ),
        ])

        self.assertEqual(structure.nodes, [])
        self.assertEqual(authority.records[0].authority_class, AUTHORITY_LOCAL_OR_UNSUPPORTED)

    def test_step_9_structure_is_immutable(self):
        items = [
            heading(1, level=1, ilvl=0),
            heading(2, level=1, ilvl=1),
            numbered(3, ilvl=0),
        ]
        before = preserve_explicit_docx_hierarchy(items, document_id="doc")
        classify(items)
        after = preserve_explicit_docx_hierarchy(items, document_id="doc")

        self.assertEqual(before, after)

    def test_step_10_report_is_immutable(self):
        items = [heading(1, level=1, ilvl=0), heading(2, level=1, ilvl=1)]
        structure = preserve_explicit_docx_hierarchy(items, document_id="doc")
        before = analyze_docx_heading_numbering_consistency(items, structure)
        observations = observe_docx_numbering_patterns(items)
        classify_docx_numbering_authority(items, structure, before, observations)
        after = analyze_docx_heading_numbering_consistency(items, structure)

        self.assertEqual(before, after)

    def test_step_11_report_is_immutable(self):
        items = [numbered(1, ilvl=0), numbered(2, ilvl=1)]
        structure = preserve_explicit_docx_hierarchy(items, document_id="doc")
        consistency = analyze_docx_heading_numbering_consistency(items, structure)
        before = observe_docx_numbering_patterns(items)
        classify_docx_numbering_authority(items, structure, consistency, before)
        after = observe_docx_numbering_patterns(items)

        self.assertEqual(before, after)

    def test_provenance_basis_fields_are_machine_readable(self):
        _structure, _consistency, _observations, authority = classify([
            heading(1, level=1, ilvl=0),
            numbered(2, ilvl=1),
        ])

        explicit = authority.records[0]
        ambiguous = authority.records[1]
        self.assertEqual(explicit.basis["step"], "heading_numbering_consistency")
        self.assertEqual(explicit.node_id, explicit.basis["node_id"])
        self.assertEqual(ambiguous.basis["step"], "numbering_observation")
        self.assertEqual(ambiguous.basis["numbering_depth"], 2)

    def test_repeated_agreement_does_not_automatically_create_trusted_numbering_system(self):
        _structure, _consistency, _observations, authority = classify([
            heading(1, level=1, ilvl=0, num_id=5),
            heading(2, level=2, ilvl=1, num_id=5),
            numbered(3, ilvl=2, num_id=5),
        ])

        self.assertEqual(
            [record.authority_class for record in authority.records],
            [
                AUTHORITY_EXPLICIT_CONFIRMED,
                AUTHORITY_EXPLICIT_CONFIRMED,
                AUTHORITY_AMBIGUOUS_NUMBERING,
            ],
        )
        self.assertEqual(authority.diagnostics["trusted_numbering_system_count"], 0)
        self.assertNotIn(
            AUTHORITY_TRUSTED_NUMBERING_SYSTEM,
            {record.authority_class for record in authority.records},
        )

    def test_adapter_backed_docx_flow_produces_authority_report(self):
        Document, _OxmlElement, _qn = _require_docx()
        document = Document()
        heading_one = document.add_heading("Introduction", level=1)
        heading_two = document.add_heading("Background", level=2)
        list_item = document.add_paragraph("Nested numbered body item")
        _set_direct_numbering(heading_one, num_id=5, ilvl=0)
        _set_direct_numbering(heading_two, num_id=5, ilvl=1)
        _set_direct_numbering(list_item, num_id=5, ilvl=2)
        buffer = io.BytesIO()
        document.save(buffer)

        canonical = adapt_docx_document(buffer.getvalue(), "authority.docx")
        evidence = map_document_to_structural_evidence(canonical)
        structure = preserve_explicit_docx_hierarchy(evidence, document_id="doc")
        consistency = analyze_docx_heading_numbering_consistency(evidence, structure)
        observations = observe_docx_numbering_patterns(evidence)
        authority = classify_docx_numbering_authority(
            evidence,
            structure,
            consistency,
            observations,
        )

        self.assertEqual(
            [record.authority_class for record in authority.records],
            [
                AUTHORITY_EXPLICIT_CONFIRMED,
                AUTHORITY_EXPLICIT_CONFIRMED,
                AUTHORITY_AMBIGUOUS_NUMBERING,
            ],
        )
        self.assertEqual([node.level for node in structure.nodes], [1, 2])

    def test_diagnostic_rendering_is_not_a_hierarchy_tree(self):
        _structure, _consistency, _observations, authority = classify([
            heading(1, level=1, ilvl=0),
            numbered(2, ilvl=1),
        ])

        self.assertEqual(render_docx_numbering_authority(authority), "\n".join([
            "[EXPLICIT_CONFIRMED] ev-0001 node=docx-heading-0001 group=num:5",
            "  heading=1 numbering=1",
            "[AMBIGUOUS_NUMBERING] ev-0002 group=num:5",
            "  depth=2 transition=start",
        ]))


if __name__ == "__main__":
    unittest.main()
