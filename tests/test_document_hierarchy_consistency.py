import io
import sys
import unittest
from pathlib import Path
from typing import Optional


BACKEND_DIR = Path(__file__).resolve().parents[1]
if str(BACKEND_DIR) not in sys.path:
    sys.path.insert(0, str(BACKEND_DIR))

from canonical_document import SourcePosition, adapt_docx_document
from document_hierarchy_consistency import (
    CONSISTENCY_CONFLICTING,
    CONSISTENCY_CONSISTENT,
    CONSISTENCY_INSUFFICIENT,
    CONSISTENCY_NOT_NUMBERED,
    analyze_docx_heading_numbering_consistency,
    render_heading_numbering_consistency,
)
from document_hierarchy_contracts import (
    CONFIDENCE_HIGH,
    CONTENT_KIND_DOCUMENT_BOUNDARY,
    CONTENT_KIND_LIST_ITEM,
    EVIDENCE_ROLE_HEADING,
    EVIDENCE_ROLE_LIST,
    SOURCE_FORMAT_DOCX,
    SourceSpan,
    StructuralEvidence,
    STATUS_UNRESOLVED,
)
from document_hierarchy_evidence_mapper import map_document_to_structural_evidence
from document_hierarchy_explicit_preserver import preserve_explicit_docx_hierarchy


def span(order: int) -> SourceSpan:
    position = SourcePosition("document_flow", 1, order, 0, 10)
    return SourceSpan(start=position, end=position)


def heading(
    text: str,
    *,
    order: int,
    level: int,
    numbering_ilvl: Optional[int] = None,
    numbering_source: Optional[str] = "direct",
    numbering=None,
) -> StructuralEvidence:
    if numbering is None:
        numbering = {}
        if numbering_ilvl is not None:
            numbering = {
                "source": numbering_source,
                "num_id": 5,
                "ilvl": numbering_ilvl,
                "list_level": numbering_ilvl,
            }
    return StructuralEvidence(
        evidence_id=f"ev-{order:04d}",
        source_format=SOURCE_FORMAT_DOCX,
        source_order=order,
        source_span=span(order),
        raw_text=text,
        normalized_text=text,
        evidence_role=EVIDENCE_ROLE_HEADING,
        content_kind=CONTENT_KIND_DOCUMENT_BOUNDARY,
        native_evidence={
            "role_hint": "heading",
            "style_hint": f"Heading {level}",
            "native_hierarchy_hint": level,
        },
        numbering_evidence=numbering,
        confidence_hint=CONFIDENCE_HIGH,
    )


def local_list(order: int, *, ilvl: int = 1) -> StructuralEvidence:
    return StructuralEvidence(
        evidence_id=f"ev-{order:04d}",
        source_format=SOURCE_FORMAT_DOCX,
        source_order=order,
        source_span=span(order),
        raw_text="Local list",
        normalized_text="Local list",
        evidence_role=EVIDENCE_ROLE_LIST,
        content_kind=CONTENT_KIND_LIST_ITEM,
        native_evidence={"role_hint": "list", "native_hierarchy_hint": ilvl},
        numbering_evidence={"source": "direct", "num_id": 5, "ilvl": ilvl},
        confidence_hint=CONFIDENCE_HIGH,
    )


def analyze(items):
    structure = preserve_explicit_docx_hierarchy(items, document_id="doc")
    return structure, analyze_docx_heading_numbering_consistency(items, structure)


def _require_docx():
    try:
        from docx import Document
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
    except ImportError as exc:  # pragma: no cover - dependency check.
        raise unittest.SkipTest("python-docx is not installed") from exc
    return Document, OxmlElement, qn


def _set_direct_numbering(paragraph, *, num_id: int = 5, ilvl: int = 0) -> None:
    _Document, OxmlElement, qn = _require_docx()
    p_pr = paragraph._p.get_or_add_pPr()  # noqa: SLF001 - test fixture XML.
    num_pr = p_pr.numPr
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        p_pr.append(num_pr)
    ilvl_el = num_pr.find(qn("w:ilvl"))
    if ilvl_el is None:
        ilvl_el = OxmlElement("w:ilvl")
        num_pr.append(ilvl_el)
    ilvl_el.set(qn("w:val"), str(ilvl))
    num_id_el = num_pr.find(qn("w:numId"))
    if num_id_el is None:
        num_id_el = OxmlElement("w:numId")
        num_pr.append(num_id_el)
    num_id_el.set(qn("w:val"), str(num_id))


class DocxHeadingNumberingConsistencyTests(unittest.TestCase):
    def test_matching_heading_1_is_consistent(self):
        _structure, report = analyze([heading("A", order=1, level=1, numbering_ilvl=0)])
        self.assertEqual(report.records[0].result, CONSISTENCY_CONSISTENT)
        self.assertEqual(report.records[0].heading_level, 1)
        self.assertEqual(report.records[0].numbering_level, 1)

    def test_matching_heading_2_is_consistent(self):
        _structure, report = analyze([
            heading("A", order=1, level=1, numbering_ilvl=0),
            heading("B", order=2, level=2, numbering_ilvl=1),
        ])
        self.assertEqual(report.records[1].result, CONSISTENCY_CONSISTENT)
        self.assertEqual(report.records[1].numbering_level, 2)

    def test_matching_heading_3_is_consistent(self):
        _structure, report = analyze([
            heading("A", order=1, level=1, numbering_ilvl=0),
            heading("B", order=2, level=2, numbering_ilvl=1),
            heading("C", order=3, level=3, numbering_ilvl=2),
        ])
        self.assertEqual(report.records[2].result, CONSISTENCY_CONSISTENT)
        self.assertEqual(report.records[2].numbering_level, 3)

    def test_conflict_does_not_mutate_canonical_level(self):
        structure, report = analyze([heading("A", order=1, level=1, numbering_ilvl=2)])
        node = structure.nodes[0]
        record = report.records[0]

        self.assertEqual(node.level, 1)
        self.assertIsNone(node.parent_id)
        self.assertEqual(record.heading_level, 1)
        self.assertEqual(record.numbering_level, 3)
        self.assertEqual(record.result, CONSISTENCY_CONFLICTING)

    def test_no_numbering_is_not_numbered_without_penalty(self):
        structure, report = analyze([heading("B", order=1, level=2)])

        self.assertEqual(structure.nodes[0].level, 2)
        self.assertEqual(report.records[0].result, CONSISTENCY_NOT_NUMBERED)
        self.assertIsNone(report.records[0].numbering_level)

    def test_numbering_without_ilvl_is_insufficient(self):
        _structure, report = analyze([
            heading("A", order=1, level=1, numbering={"source": "style_linked", "num_id": 5})
        ])

        self.assertEqual(report.records[0].result, CONSISTENCY_INSUFFICIENT)
        self.assertEqual(report.records[0].numbering_source, "style_linked")

    def test_local_numbered_list_is_ignored(self):
        structure, report = analyze([
            heading("A", order=1, level=1, numbering_ilvl=0),
            local_list(2, ilvl=1),
        ])

        self.assertEqual(len(structure.nodes), 1)
        self.assertEqual(len(report.records), 1)
        self.assertEqual(report.diagnostics["explicit_heading_count"], 1)

    def test_all_heading_1_with_varying_numbering_levels_keeps_step_9_hierarchy(self):
        structure, report = analyze([
            heading("A", order=1, level=1, numbering_ilvl=0),
            heading("B", order=2, level=1, numbering_ilvl=1),
            heading("C", order=3, level=1, numbering_ilvl=2),
        ])

        self.assertEqual([node.level for node in structure.nodes], [1, 1, 1])
        self.assertEqual([node.parent_id for node in structure.nodes], [None, None, None])
        self.assertEqual([record.result for record in report.records], [
            CONSISTENCY_CONSISTENT,
            CONSISTENCY_CONFLICTING,
            CONSISTENCY_CONFLICTING,
        ])

    def test_malformed_hierarchy_is_not_repaired_by_numbering(self):
        structure, report = analyze([
            heading("A", order=1, level=1, numbering_ilvl=0),
            heading("C", order=2, level=3, numbering_ilvl=1),
        ])

        node_c = structure.nodes[1]
        record_c = report.records[1]
        self.assertEqual(node_c.level, 3)
        self.assertIsNone(node_c.parent_id)
        self.assertEqual(node_c.status, STATUS_UNRESOLVED)
        self.assertEqual(record_c.result, CONSISTENCY_CONFLICTING)
        self.assertEqual(record_c.numbering_level, 2)

    def test_provenance_links_consistency_to_node_and_evidence(self):
        structure, report = analyze([heading("A", order=1, level=1, numbering_ilvl=0, numbering_source="style")])
        record = report.records[0]

        self.assertEqual(record.evidence_id, "ev-0001")
        self.assertEqual(record.node_id, structure.nodes[0].node_id)
        self.assertEqual(record.numbering_source, "style")

    def test_summary_counts(self):
        _structure, report = analyze([
            heading("A", order=1, level=1, numbering_ilvl=0),
            heading("B", order=2, level=2, numbering_ilvl=2),
            heading("C", order=3, level=3),
            heading("D", order=4, level=3, numbering={"source": "style"}),
        ])

        self.assertEqual(report.diagnostics, {
            "explicit_heading_count": 4,
            "numbered_heading_count": 2,
            "consistent_heading_count": 1,
            "conflicting_heading_count": 1,
            "uncomparable_heading_count": 2,
        })

    def test_rendering_is_diagnostic_only(self):
        structure, report = analyze([heading("A", order=1, level=1, numbering_ilvl=0)])
        self.assertEqual(render_heading_numbering_consistency(structure, report), "\n".join([
            "L1 A",
            "  numbering: numbering_level=1 -> consistent",
        ]))

    def test_adapter_backed_docx_flow_uses_genuine_word_numbering_metadata(self):
        Document, _OxmlElement, _qn = _require_docx()
        document = Document()
        heading_one = document.add_heading("Introduction", level=1)
        heading_two = document.add_heading("Background", level=2)
        _set_direct_numbering(heading_one, num_id=5, ilvl=0)
        _set_direct_numbering(heading_two, num_id=5, ilvl=1)
        buffer = io.BytesIO()
        document.save(buffer)

        canonical = adapt_docx_document(buffer.getvalue(), "numbered-headings.docx")
        evidence = map_document_to_structural_evidence(canonical)
        structure = preserve_explicit_docx_hierarchy(evidence, document_id="doc")
        report = analyze_docx_heading_numbering_consistency(evidence, structure)

        self.assertEqual([record.result for record in report.records], [
            CONSISTENCY_CONSISTENT,
            CONSISTENCY_CONSISTENT,
        ])
        self.assertEqual([record.numbering_source for record in report.records], ["direct", "direct"])


if __name__ == "__main__":
    unittest.main()
