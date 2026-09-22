import io
import json
import sys
import types
import unittest
from pathlib import Path


BACKEND_DIR = Path(__file__).resolve().parents[1]
if str(BACKEND_DIR) not in sys.path:
    sys.path.insert(0, str(BACKEND_DIR))

from canonical_document import (
    ROLE_BODY,
    ROLE_HEADING,
    ROLE_LIST,
    ROLE_TABLE,
    ROLE_TITLE,
    CanonicalBlock,
    CanonicalDocumentInput,
    CanonicalUnit,
    SourcePosition,
    adapt_docx_document,
    adapt_pdf_document,
    adapt_pptx_document,
)
from universal_structure_interpreter import (
    DOCUMENT_STRUCTURE_SCHEMA,
    UniversalStructureError,
    build_response_input,
    build_structured_text_representation,
    interpret_document_structure,
    parse_model_response,
    structure_from_model_json,
)
from universal_structure_mapping import (
    STATUS_ACCEPTED,
    STATUS_FALLBACK,
    STATUS_PARTIAL,
    anchor_structure,
    build_mapping_result,
    build_section_path,
    diagnostic_chunks_from_segments,
    evaluate_mapping,
    production_compatible_chunk_text,
    segment_document,
)


def block(unit_type, unit_index, block_index, text, role=ROLE_BODY, style=None, hierarchy=None, order=None):
    order = order if order is not None else (unit_index * 100) + block_index
    return CanonicalBlock(
        block_id=f"{unit_type}-{unit_index}-{block_index}",
        text=text,
        raw_text=text,
        source_order=order,
        source_position=SourcePosition(unit_type, unit_index, block_index, 0, len(text)),
        role_hint=role,
        style_hint=style,
        native_hierarchy_hint=hierarchy,
    )


def doc(blocks, file_type="pdf", unit_type="page"):
    units_by_index = {}
    for item in blocks:
        units_by_index.setdefault(item.source_position.unit_index, []).append(item)
    units = [
        CanonicalUnit(
            unit_id=f"{unit_type}-{index}",
            unit_type=unit_type,
            unit_index=index,
            display_label=f"{unit_type.title()} {index}",
            blocks=unit_blocks,
        )
        for index, unit_blocks in sorted(units_by_index.items())
    ]
    return CanonicalDocumentInput(file_type, f"fixture.{file_type}", "Fixture", units)


def model(sections, confidence="HIGH"):
    return structure_from_model_json(
        {
            "document_title": "Fixture",
            "document_type": "notes",
            "structure_confidence": confidence,
            "sections": sections,
            "local_structure_summary": {
                "detected": False,
                "types": [],
                "representative_examples": [],
            },
            "notes": [],
        }
    )


def section(title, level=1, parent=None, unit_type="page", unit=1, block_index=1, confidence="HIGH"):
    semantic_title = title.replace("1 ", "").replace("2 ", "") if title else "Model Section"
    return {
        "source_title": title,
        "semantic_title": semantic_title,
        "level": level,
        "parent": parent,
        "expected_start": {
            "unit_type": unit_type,
            "unit_index": unit,
            "block_index": block_index,
            "start_offset": 0,
            "end_offset": None,
        },
        "expected_end": {
            "unit_type": unit_type,
            "unit_index": unit,
            "block_index": block_index,
            "start_offset": 0,
            "end_offset": None,
        },
        "confidence": confidence,
        "scope": "global",
    }


def assert_strict_schema_objects(test_case, schema, path="schema"):
    schema_type = schema.get("type")
    schema_types = schema_type if isinstance(schema_type, list) else [schema_type]
    if "object" in schema_types:
        properties = schema.get("properties", {})
        test_case.assertIsInstance(properties, dict, path)
        test_case.assertEqual(
            set(schema.get("required", [])),
            set(properties.keys()),
            f"{path} required keys must match properties",
        )
        test_case.assertFalse(
            schema.get("additionalProperties", True),
            f"{path} must disable additionalProperties",
        )
    for key, subschema in schema.get("properties", {}).items():
        if isinstance(subschema, dict):
            assert_strict_schema_objects(test_case, subschema, f"{path}.properties.{key}")
    items = schema.get("items")
    if isinstance(items, dict):
        assert_strict_schema_objects(test_case, items, f"{path}.items")
    for keyword in ("anyOf", "oneOf", "allOf"):
        for index, subschema in enumerate(schema.get(keyword, []) or []):
            if isinstance(subschema, dict):
                assert_strict_schema_objects(test_case, subschema, f"{path}.{keyword}[{index}]")


class FakeResponses:
    def __init__(self, response=None, error=None):
        self.response = response
        self.error = error
        self.calls = []

    def create(self, **kwargs):
        self.calls.append(kwargs)
        if self.error:
            raise self.error
        return self.response


class FakeClient:
    def __init__(self, response=None, error=None):
        self.responses = FakeResponses(response, error)


class UniversalStructureMappingTests(unittest.TestCase):
    def test_01_canonical_document_input_creation(self):
        document = doc([block("page", 1, 1, "1 Intro")])
        self.assertEqual(document.file_type, "pdf")
        self.assertEqual(len(document.units), 1)
        self.assertEqual(document.blocks[0].source_position.unit_type, "page")

    def test_02_source_position_ordering(self):
        first = SourcePosition("page", 1, 1, 0, 5)
        second = SourcePosition("page", 1, 1, 10, 20)
        third = SourcePosition("page", 2, 1, 0, 5)
        self.assertLess(first, second)
        self.assertLess(second, third)

    def test_03_exact_anchoring(self):
        document = doc([block("page", 1, 1, "1 Introduction\nBody")])
        anchored = anchor_structure(document, model([section("1 Introduction")]))
        self.assertEqual(anchored.anchor_summary.exact_count, 1)

    def test_04_normalized_anchoring(self):
        document = doc([block("page", 1, 1, "1   Introduction\nBody")])
        anchored = anchor_structure(document, model([section("1 Introduction")]))
        self.assertEqual(anchored.anchor_summary.normalized_count, 1)

    def test_05_fuzzy_anchoring(self):
        document = doc([block("page", 1, 1, "1 Prokaryotic gene regulation: operons\nBody")])
        anchored = anchor_structure(document, model([section("1 Prokaryotic regulation operons")]))
        self.assertEqual(anchored.anchor_summary.fuzzy_count, 1)

    def test_06_unresolved_heading(self):
        document = doc([block("page", 1, 1, "Body only")])
        anchored = anchor_structure(document, model([section("1 Missing")]))
        self.assertEqual(anchored.anchor_summary.unresolved, 1)

    def test_07_duplicate_collision(self):
        document = doc([block("page", 1, 1, "1 Repeat\nBody\n1 Repeat\nMore")])
        anchored = anchor_structure(document, model([section("1 Repeat"), section("1 Repeat")]))
        self.assertEqual(anchored.anchor_summary.duplicate_anchor_collisions, 1)
        self.assertEqual(anchored.anchor_summary.anchored, 1)

    def test_08_source_order_conflict(self):
        document = doc([block("page", 1, 1, "2 Later\nBody\n1 Earlier\nBody")])
        anchored = anchor_structure(document, model([section("1 Earlier"), section("2 Later")]))
        self.assertEqual(anchored.anchor_summary.source_order_violations, 1)

    def test_09_nested_hierarchy_path(self):
        document = doc([block("page", 1, 1, "1 Parent\nBody\n1.1 Child\nChild body")])
        structure = model([
            section("1 Parent"),
            section("1.1 Child", level=2, parent="1 Parent"),
        ])
        anchored = anchor_structure(document, structure)
        sections_by_id = {item.section_id: item for item in anchored.sections}
        self.assertEqual(build_section_path("section-0002", sections_by_id), "1 Parent > 1.1 Child")

    def test_10_same_unit_multiple_headings(self):
        raw = "Preface\n1 One\nA\n1.1 Child\nB\n2 Two\nC"
        document = doc([block("page", 1, 1, raw)])
        anchored = anchor_structure(document, model([
            section("1 One"),
            section("1.1 Child", level=2, parent="1 One"),
            section("2 Two"),
        ]))
        segments = segment_document(document, anchored)
        self.assertEqual([s.section_path for s in segments][1:], ["1 One", "1 One > 1.1 Child", "2 Two"])

    def test_11_continuation_across_units(self):
        document = doc([
            block("page", 1, 1, "1 Intro\nPage one"),
            block("page", 2, 1, "Continuation page two"),
        ])
        anchored = anchor_structure(document, model([section("1 Intro")]))
        segments = segment_document(document, anchored)
        self.assertEqual(segments[-1].section_path, "1 Intro")

    def test_12_partial_mapping(self):
        document = doc([block("page", 1, 1, "1 One\nBody")])
        result = build_mapping_result(document, model([section("1 One"), section("2 Missing", unit=2)]))
        self.assertEqual(result.mapping_status, STATUS_PARTIAL)

    def test_13_fallback_mapping_for_low_confidence(self):
        document = doc([block("page", 1, 1, "1 One\nBody")])
        result = build_mapping_result(document, model([section("1 One", confidence="LOW")], confidence="LOW"))
        self.assertEqual(result.mapping_status, STATUS_FALLBACK)

    def test_14_model_only_semantic_parent_can_still_allow_partial_children(self):
        document = doc([block("page", 1, 1, "1.1 Child\nBody")])
        structure = model([
            {
                **section(None),
                "source_title": None,
                "semantic_title": "Model Parent",
                "expected_start": None,
                "expected_end": None,
            },
            section("1.1 Child", level=2, parent="Model Parent"),
        ])
        result = build_mapping_result(document, structure)
        self.assertEqual(result.sections_anchored, 1)
        self.assertEqual(result.mapping_status, STATUS_PARTIAL)

    def test_15_no_segment_boundary_crossing(self):
        raw = "1 One\n" + ("Alpha " * 40) + "\n2 Two\n" + ("Beta " * 40)
        document = doc([block("page", 1, 1, raw)])
        anchored = anchor_structure(document, model([section("1 One"), section("2 Two")]))
        segments = segment_document(document, anchored)
        self.assertEqual(sum(1 for segment in segments if "Alpha" in segment.text and "Beta" in segment.text), 0)

    def test_16_no_chunk_boundary_crossing(self):
        raw = "1 One\n" + ("Alpha " * 120) + "\n2 Two\n" + ("Beta " * 120)
        document = doc([block("page", 1, 1, raw)])
        chunks = diagnostic_chunks_from_segments(segment_document(document, anchor_structure(document, model([section("1 One"), section("2 Two")]))))
        for chunk in chunks:
            self.assertFalse("Alpha" in chunk.text and "Beta" in chunk.text)

    def test_17_pdf_multiple_headings_on_one_page(self):
        self.test_10_same_unit_multiple_headings()

    def test_18_pdf_page_continuation(self):
        self.test_11_continuation_across_units()

    def test_19_docx_heading_style_preserved_as_hint(self):
        from docx import Document
        document = Document()
        document.add_heading("Chapter One", level=1)
        document.add_paragraph("Body")
        buffer = io.BytesIO()
        document.save(buffer)
        canonical = adapt_docx_document(buffer.getvalue(), "sample.docx")
        self.assertEqual(canonical.blocks[0].role_hint, ROLE_HEADING)
        self.assertEqual(canonical.blocks[0].native_hierarchy_hint, 1)

    def test_20_docx_wrong_heading_style_does_not_force_hierarchy(self):
        document = doc([block("document_flow", 1, 1, "Styled body", role=ROLE_HEADING, style="Heading 1", hierarchy=1)], file_type="docx", unit_type="document_flow")
        result = build_mapping_result(document, model([section("Not in source", unit_type="document_flow")]))
        self.assertEqual(result.mapping_status, STATUS_FALLBACK)

    def test_21_docx_multiple_headings_in_consecutive_paragraphs(self):
        document = doc([
            block("document_flow", 1, 1, "1 First", role=ROLE_HEADING),
            block("document_flow", 1, 2, "1.1 Second", role=ROLE_HEADING),
        ], file_type="docx", unit_type="document_flow")
        result = build_mapping_result(document, model([
            section("1 First", unit_type="document_flow", block_index=1),
            section("1.1 Second", level=2, parent="1 First", unit_type="document_flow", block_index=2),
        ]))
        self.assertEqual(result.mapping_status, STATUS_ACCEPTED)

    def test_22_docx_list_hierarchy_preserved_as_hint(self):
        document = doc([block("document_flow", 1, 1, "List item", role=ROLE_LIST, hierarchy=2)], file_type="docx", unit_type="document_flow")
        self.assertEqual(document.blocks[0].role_hint, ROLE_LIST)
        self.assertEqual(document.blocks[0].native_hierarchy_hint, 2)

    def test_23_docx_table_block_ordering(self):
        document = doc([
            block("document_flow", 1, 1, "Heading", role=ROLE_HEADING, order=1),
            block("document_flow", 1, 2, "A | B", role=ROLE_TABLE, order=2),
        ], file_type="docx", unit_type="document_flow")
        self.assertEqual([b.role_hint for b in document.blocks], [ROLE_HEADING, ROLE_TABLE])

    def test_24_pptx_slide_title_preserved(self):
        try:
            from pptx import Presentation
        except ImportError:
            self.skipTest("python-pptx is not installed in this local venv")
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[0])
        slide.shapes.title.text = "Cardiac Muscle"
        buffer = io.BytesIO()
        presentation.save(buffer)
        canonical = adapt_pptx_document(buffer.getvalue(), "sample.pptx")
        self.assertEqual(canonical.units[0].display_label, "Cardiac Muscle")
        self.assertEqual(canonical.blocks[0].role_hint, ROLE_TITLE)

    def test_25_pptx_repeated_slide_titles(self):
        document = doc([
            block("slide", 1, 1, "Action potential", role=ROLE_TITLE),
            block("slide", 2, 1, "Action potential", role=ROLE_TITLE),
        ], file_type="pptx", unit_type="slide")
        self.assertEqual([u.blocks[0].text for u in document.units], ["Action potential", "Action potential"])

    def test_26_pptx_multiple_text_shapes_same_slide(self):
        document = doc([
            block("slide", 1, 1, "Title", role=ROLE_TITLE),
            block("slide", 1, 2, "Body A", role=ROLE_BODY),
            block("slide", 1, 3, "Body B", role=ROLE_BODY),
        ], file_type="pptx", unit_type="slide")
        self.assertEqual(len(document.blocks), 3)

    def test_27_pptx_bullet_depth_preserved(self):
        document = doc([block("slide", 1, 2, "Nested bullet", role=ROLE_LIST, hierarchy=1)], file_type="pptx", unit_type="slide")
        self.assertEqual(document.blocks[0].native_hierarchy_hint, 1)

    def test_28_pptx_title_placeholder_is_hint_not_truth(self):
        document = doc([block("slide", 1, 1, "Decorative title", role=ROLE_TITLE)], file_type="pptx", unit_type="slide")
        result = build_mapping_result(document, model([section("Missing semantic section", unit_type="slide")]))
        self.assertEqual(result.mapping_status, STATUS_FALLBACK)

    def test_29_pptx_section_spanning_multiple_slides(self):
        document = doc([
            block("slide", 1, 1, "1 Cardiac Muscle", role=ROLE_TITLE),
            block("slide", 2, 1, "Continuation", role=ROLE_BODY),
        ], file_type="pptx", unit_type="slide")
        result = build_mapping_result(document, model([section("1 Cardiac Muscle", unit_type="slide")]))
        self.assertEqual(result.segments[-1].section_path, "1 Cardiac Muscle")

    def test_30_mocked_pdf_model_call(self):
        response = fake_response()
        client = FakeClient(response=response)
        document = doc([block("page", 1, 1, "1 Intro")])
        interpret_document_structure(document=document, model="mock", client=client, original_file_bytes=b"%PDF")
        content = client.responses.calls[0]["input"][0]["content"]
        self.assertEqual(content[1]["type"], "input_file")

    def test_31_mocked_structured_docx_model_input(self):
        document = doc([block("document_flow", 1, 1, "1 Intro", role=ROLE_HEADING)], file_type="docx", unit_type="document_flow")
        payload = build_response_input(document)
        self.assertEqual(payload[0]["content"][1]["type"], "input_text")
        self.assertIn("role_hint=heading", payload[0]["content"][1]["text"])

    def test_32_mocked_structured_pptx_model_input(self):
        document = doc([block("slide", 1, 1, "Title", role=ROLE_TITLE)], file_type="pptx", unit_type="slide")
        text = build_structured_text_representation(document)
        self.assertIn("[unit=slide:1", text)
        self.assertIn("role_hint=title", text)

    def test_33_malformed_model_output(self):
        with self.assertRaises(UniversalStructureError):
            parse_model_response("not json")

    def test_34_incomplete_model_response(self):
        diagnostics = types.SimpleNamespace(response_status="incomplete", incomplete_reason="max_output_tokens")
        with self.assertRaises(UniversalStructureError):
            parse_model_response("{", diagnostics)

    def test_35_low_confidence_structure(self):
        document = doc([block("page", 1, 1, "1 Intro")])
        result = build_mapping_result(document, model([section("1 Intro")], confidence="LOW"))
        self.assertEqual(result.structure_status, STATUS_FALLBACK)

    def test_existing_chunker_adapter_shape(self):
        chunks = production_compatible_chunk_text("A" * 1200, max_chars=1000, overlap=200)
        self.assertEqual(len(chunks), 2)

    def test_strict_schema_expected_start_requires_all_source_position_properties(self):
        position_schema = DOCUMENT_STRUCTURE_SCHEMA["properties"]["sections"]["items"]["properties"]["expected_start"]
        self.assertEqual(
            set(position_schema["required"]),
            {"unit_type", "unit_index", "block_index", "start_offset", "end_offset"},
        )

    def test_strict_schema_expected_end_requires_all_source_position_properties(self):
        position_schema = DOCUMENT_STRUCTURE_SCHEMA["properties"]["sections"]["items"]["properties"]["expected_end"]
        self.assertEqual(
            set(position_schema["required"]),
            {"unit_type", "unit_index", "block_index", "start_offset", "end_offset"},
        )

    def test_source_position_schema_keeps_nullable_unknown_fields(self):
        position_schema = DOCUMENT_STRUCTURE_SCHEMA["properties"]["sections"]["items"]["properties"]["expected_start"]
        self.assertIn("null", position_schema["type"])
        for property_name in ("unit_type", "unit_index", "block_index", "start_offset", "end_offset"):
            self.assertIn("null", position_schema["properties"][property_name]["type"])

    def test_all_nested_object_schemas_are_strict(self):
        assert_strict_schema_objects(self, DOCUMENT_STRUCTURE_SCHEMA)

    def test_optional_semantic_values_can_be_null(self):
        section_schema = DOCUMENT_STRUCTURE_SCHEMA["properties"]["sections"]["items"]["properties"]
        self.assertIn("null", section_schema["source_title"]["type"])
        self.assertIn("null", section_schema["parent"]["type"])
        self.assertIn("null", section_schema["expected_start"]["type"])
        self.assertIn("null", section_schema["expected_end"]["type"])

    def test_model_response_with_null_positions_parses(self):
        parsed = model([
            {
                **section("1 Intro"),
                "expected_start": {
                    "unit_type": "page",
                    "unit_index": 1,
                    "block_index": None,
                    "start_offset": None,
                    "end_offset": None,
                },
                "expected_end": None,
            }
        ])
        self.assertIsNone(parsed.sections[0].expected_start)
        self.assertIsNone(parsed.sections[0].expected_end)


def fake_response():
    raw = json.dumps(
        {
            "document_title": "Fixture",
            "document_type": "notes",
            "structure_confidence": "HIGH",
            "sections": [
                section("1 Intro"),
            ],
            "local_structure_summary": {
                "detected": False,
                "types": [],
                "representative_examples": [],
            },
            "notes": [],
        }
    )
    return types.SimpleNamespace(
        output_text=raw,
        status="completed",
        incomplete_details=None,
        max_output_tokens=6000,
        output=[types.SimpleNamespace(status="completed", content=[])],
        usage=types.SimpleNamespace(
            input_tokens=10,
            output_tokens=20,
            total_tokens=30,
            output_tokens_details=types.SimpleNamespace(reasoning_tokens=0),
        ),
    )


if __name__ == "__main__":
    unittest.main()
