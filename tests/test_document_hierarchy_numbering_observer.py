import io
import sys
import unittest
from pathlib import Path
from typing import Optional


BACKEND_DIR = Path(__file__).resolve().parents[1]
if str(BACKEND_DIR) not in sys.path:
    sys.path.insert(0, str(BACKEND_DIR))

from canonical_document import SourcePosition, adapt_docx_document
from document_hierarchy_consistency import analyze_docx_heading_numbering_consistency
from document_hierarchy_contracts import (
    CONFIDENCE_HIGH,
    CONTENT_KIND_BODY,
    CONTENT_KIND_DOCUMENT_BOUNDARY,
    CONTENT_KIND_LIST_ITEM,
    EVIDENCE_ROLE_BODY,
    EVIDENCE_ROLE_HEADING,
    EVIDENCE_ROLE_LIST,
    SOURCE_FORMAT_DOCX,
    SourceSpan,
    StructuralEvidence,
)
from document_hierarchy_evidence_mapper import map_document_to_structural_evidence
from document_hierarchy_explicit_preserver import preserve_explicit_docx_hierarchy
from document_hierarchy_numbering_observer import (
    TRANSITION_ASCEND,
    TRANSITION_DESCEND,
    TRANSITION_INSUFFICIENT,
    TRANSITION_RESTART,
    TRANSITION_SAME_DEPTH,
    TRANSITION_START,
    observe_docx_numbering_patterns,
    render_docx_numbering_observations,
)


def span(order: int) -> SourceSpan:
    position = SourcePosition("document_flow", 1, order, 0, 10)
    return SourceSpan(start=position, end=position)


def numbered(
    order: int,
    *,
    ilvl: Optional[int],
    num_id=5,
    abstract_num_id=7,
    role=EVIDENCE_ROLE_LIST,
    source="direct",
    start=None,
    start_override=None,
) -> StructuralEvidence:
    numbering = {
        "source": source,
        "num_id": num_id,
        "abstract_num_id": abstract_num_id,
        "format": "decimal",
        "level_text": "%1.",
    }
    if ilvl is not None:
        numbering["ilvl"] = ilvl
        numbering["list_level"] = ilvl
    if start is not None:
        numbering["start"] = start
    if start_override is not None:
        numbering["start_override"] = start_override
    return StructuralEvidence(
        evidence_id=f"ev-{order:04d}",
        source_format=SOURCE_FORMAT_DOCX,
        source_order=order,
        source_span=span(order),
        raw_text=f"Numbered {order}",
        normalized_text=f"Numbered {order}",
        evidence_role=role,
        content_kind=CONTENT_KIND_LIST_ITEM,
        native_evidence={"role_hint": "list", "native_hierarchy_hint": ilvl},
        numbering_evidence=numbering,
        confidence_hint=CONFIDENCE_HIGH,
    )


def body(order: int) -> StructuralEvidence:
    return StructuralEvidence(
        evidence_id=f"ev-{order:04d}",
        source_format=SOURCE_FORMAT_DOCX,
        source_order=order,
        source_span=span(order),
        raw_text=f"Body {order}",
        normalized_text=f"Body {order}",
        evidence_role=EVIDENCE_ROLE_BODY,
        content_kind=CONTENT_KIND_BODY,
        native_evidence={"role_hint": "body"},
    )


def heading(order: int, *, level: int, ilvl: Optional[int] = None) -> StructuralEvidence:
    numbering = {}
    if ilvl is not None:
        numbering = {"source": "direct", "num_id": 5, "abstract_num_id": 7, "ilvl": ilvl}
    return StructuralEvidence(
        evidence_id=f"ev-{order:04d}",
        source_format=SOURCE_FORMAT_DOCX,
        source_order=order,
        source_span=span(order),
        raw_text=f"Heading {order}",
        normalized_text=f"Heading {order}",
        evidence_role=EVIDENCE_ROLE_HEADING,
        content_kind=CONTENT_KIND_DOCUMENT_BOUNDARY,
        native_evidence={
            "role_hint": "heading",
            "style_hint": f"Heading {level}",
            "native_hierarchy_hint": level,
        },
        numbering_evidence=numbering,
        confidence_hint=CONFIDENCE_HIGH,
    )


def _require_docx():
    try:
        from docx import Document
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
    except ImportError as exc:  # pragma: no cover - dependency check.
        raise unittest.SkipTest("python-docx is not installed") from exc
    return Document, OxmlElement, qn


def _set_direct_numbering(paragraph, *, num_id: int = 5, ilvl: int = 0) -> None:
    _Document, OxmlElement, qn = _require_docx()
    p_pr = paragraph._p.get_or_add_pPr()  # noqa: SLF001 - test fixture XML.
    num_pr = p_pr.numPr
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        p_pr.append(num_pr)
    ilvl_el = num_pr.find(qn("w:ilvl"))
    if ilvl_el is None:
        ilvl_el = OxmlElement("w:ilvl")
        num_pr.append(ilvl_el)
    ilvl_el.set(qn("w:val"), str(ilvl))
    num_id_el = num_pr.find(qn("w:numId"))
    if num_id_el is None:
        num_id_el = OxmlElement("w:numId")
        num_pr.append(num_id_el)
    num_id_el.set(qn("w:val"), str(num_id))


class DocxNumberingObserverTests(unittest.TestCase):
    def test_flat_native_numbering_observes_same_group_and_depth_one(self):
        report = observe_docx_numbering_patterns([
            numbered(1, ilvl=0),
            numbered(2, ilvl=0),
            numbered(3, ilvl=0),
        ])

        self.assertEqual([obs.group_id for obs in report.observations], ["num:5", "num:5", "num:5"])
        self.assertEqual([obs.numbering_depth for obs in report.observations], [1, 1, 1])
        self.assertEqual([obs.transition_from_previous for obs in report.observations], [
            TRANSITION_START,
            TRANSITION_SAME_DEPTH,
            TRANSITION_SAME_DEPTH,
        ])
        self.assertEqual(report.diagnostics["numbered_non_heading_count"], 3)
        self.assertEqual(report.diagnostics["numbering_group_count"], 1)

    def test_nested_numbering_preserves_depth_transitions_without_hierarchy(self):
        report = observe_docx_numbering_patterns([
            numbered(1, ilvl=0),
            numbered(2, ilvl=1),
            numbered(3, ilvl=2),
            numbered(4, ilvl=1),
            numbered(5, ilvl=0),
        ])

        self.assertEqual([obs.numbering_depth for obs in report.observations], [1, 2, 3, 2, 1])
        self.assertEqual([obs.transition_from_previous for obs in report.observations], [
            TRANSITION_START,
            TRANSITION_DESCEND,
            TRANSITION_DESCEND,
            TRANSITION_ASCEND,
            TRANSITION_ASCEND,
        ])
        self.assertEqual(report.diagnostics["depth_3_plus_count"], 1)

    def test_separate_num_id_values_are_not_merged(self):
        report = observe_docx_numbering_patterns([
            numbered(1, ilvl=0, num_id=5, abstract_num_id=7),
            numbered(2, ilvl=0, num_id=6, abstract_num_id=8),
        ])

        self.assertEqual([obs.group_id for obs in report.observations], ["num:5", "num:6"])
        self.assertEqual(report.diagnostics["numbering_group_count"], 2)

    def test_same_abstract_num_id_different_num_id_preserves_both_facts(self):
        report = observe_docx_numbering_patterns([
            numbered(1, ilvl=0, num_id=5, abstract_num_id=7),
            numbered(2, ilvl=0, num_id=6, abstract_num_id=7),
        ])

        self.assertEqual([obs.group_id for obs in report.observations], ["num:5", "num:6"])
        self.assertEqual([obs.abstract_num_id for obs in report.observations], [7, 7])
        self.assertEqual([obs.num_id for obs in report.observations], [5, 6])

    def test_interruption_by_body_paragraph_is_observable(self):
        report = observe_docx_numbering_patterns([
            numbered(1, ilvl=0),
            numbered(2, ilvl=0),
            body(3),
            numbered(4, ilvl=0),
        ])

        self.assertEqual([obs.interrupted_since_previous for obs in report.observations], [False, False, True])
        self.assertEqual(report.diagnostics["interrupted_group_count"], 1)

    def test_explicit_heading_context_does_not_create_parent_relationship(self):
        items = [
            heading(1, level=1, ilvl=0),
            numbered(2, ilvl=0),
            numbered(3, ilvl=1),
        ]
        structure = preserve_explicit_docx_hierarchy(items, document_id="doc")
        report = observe_docx_numbering_patterns(items)

        self.assertEqual(len(structure.nodes), 1)
        self.assertEqual([obs.context_heading_evidence_id for obs in report.observations], ["ev-0001", "ev-0001"])
        self.assertFalse(hasattr(report.observations[0], "parent_id"))

    def test_local_numbered_list_remains_observation_only(self):
        items = [
            body(1),
            numbered(2, ilvl=0),
            numbered(3, ilvl=0),
            numbered(4, ilvl=0),
            body(5),
        ]
        structure = preserve_explicit_docx_hierarchy(items, document_id="doc")
        report = observe_docx_numbering_patterns(items)

        self.assertEqual(structure.nodes, [])
        self.assertEqual(len(report.observations), 3)
        self.assertEqual({obs.group_id for obs in report.observations}, {"num:5"})

    def test_numbering_evidence_missing_ilvl_is_partial_observation(self):
        report = observe_docx_numbering_patterns([
            numbered(1, ilvl=None),
            numbered(2, ilvl=1),
        ])

        self.assertIsNone(report.observations[0].ilvl)
        self.assertIsNone(report.observations[0].numbering_depth)
        self.assertEqual(report.observations[1].transition_from_previous, TRANSITION_INSUFFICIENT)

    def test_restart_transition_is_observable_from_start_override(self):
        report = observe_docx_numbering_patterns([
            numbered(1, ilvl=0),
            numbered(2, ilvl=0, start_override=1),
        ])

        self.assertEqual(report.observations[1].transition_from_previous, TRANSITION_RESTART)
        self.assertEqual(report.diagnostics["restart_count"], 1)

    def test_provenance_retains_evidence_order_and_native_identity(self):
        report = observe_docx_numbering_patterns([
            numbered(7, ilvl=2, num_id=9, abstract_num_id=11, source="style_linked"),
        ])
        obs = report.observations[0]

        self.assertEqual(obs.evidence_id, "ev-0007")
        self.assertEqual(obs.source_order, 7)
        self.assertEqual(obs.num_id, 9)
        self.assertEqual(obs.abstract_num_id, 11)
        self.assertEqual(obs.numbering_source, "style_linked")

    def test_step_9_structure_is_unchanged_by_step_11_analysis(self):
        items = [
            heading(1, level=1),
            numbered(2, ilvl=0),
            numbered(3, ilvl=1),
        ]
        before = preserve_explicit_docx_hierarchy(items, document_id="doc")
        observe_docx_numbering_patterns(items)
        after = preserve_explicit_docx_hierarchy(items, document_id="doc")

        self.assertEqual(before, after)

    def test_step_10_heading_consistency_is_not_duplicated_or_mutated(self):
        items = [
            heading(1, level=1, ilvl=0),
            numbered(2, ilvl=1),
        ]
        structure = preserve_explicit_docx_hierarchy(items, document_id="doc")
        consistency_before = analyze_docx_heading_numbering_consistency(items, structure)
        numbering_report = observe_docx_numbering_patterns(items)
        consistency_after = analyze_docx_heading_numbering_consistency(items, structure)

        self.assertEqual(consistency_before, consistency_after)
        self.assertEqual(len(consistency_before.records), 1)
        self.assertEqual(len(numbering_report.observations), 1)
        self.assertEqual(numbering_report.observations[0].evidence_id, "ev-0002")

    def test_adapter_backed_docx_flow_preserves_native_numbering_observations(self):
        Document, _OxmlElement, _qn = _require_docx()
        document = Document()
        first = document.add_paragraph("First native numbered paragraph")
        second = document.add_paragraph("Nested native numbered paragraph")
        _set_direct_numbering(first, num_id=5, ilvl=0)
        _set_direct_numbering(second, num_id=5, ilvl=1)
        buffer = io.BytesIO()
        document.save(buffer)

        canonical = adapt_docx_document(buffer.getvalue(), "numbered.docx")
        evidence = map_document_to_structural_evidence(canonical)
        report = observe_docx_numbering_patterns(evidence)

        self.assertEqual([obs.numbering_depth for obs in report.observations], [1, 2])
        self.assertEqual([obs.numbering_source for obs in report.observations], ["direct", "direct"])
        self.assertEqual([obs.group_id for obs in report.observations], ["num:5", "num:5"])

    def test_diagnostic_rendering_is_not_a_document_tree(self):
        report = observe_docx_numbering_patterns([
            numbered(1, ilvl=0),
            numbered(2, ilvl=1),
        ])

        self.assertEqual(render_docx_numbering_observations(report), "\n".join([
            "NUMBERING GROUP num:5",
            "  [depth 1] ev-0001 transition=start",
            "  [depth 2] ev-0002 transition=descend",
        ]))


if __name__ == "__main__":
    unittest.main()
