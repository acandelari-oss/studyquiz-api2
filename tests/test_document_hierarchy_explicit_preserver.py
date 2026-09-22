import io
import sys
import unittest
from pathlib import Path
from typing import Optional


BACKEND_DIR = Path(__file__).resolve().parents[1]
if str(BACKEND_DIR) not in sys.path:
    sys.path.insert(0, str(BACKEND_DIR))

from canonical_document import ROLE_BODY, ROLE_HEADING, ROLE_LIST, ROLE_TABLE, SourcePosition, adapt_docx_document
from document_hierarchy_contracts import (
    CONFIDENCE_HIGH,
    CONFIDENCE_LOW,
    CONFIDENCE_MEDIUM,
    CONTENT_KIND_BODY,
    CONTENT_KIND_DOCUMENT_BOUNDARY,
    CONTENT_KIND_LIST_ITEM,
    CONTENT_KIND_TABLE,
    EVIDENCE_ROLE_BODY,
    EVIDENCE_ROLE_HEADING,
    EVIDENCE_ROLE_LIST,
    EVIDENCE_ROLE_TABLE,
    ORIGIN_EXPLICIT_SOURCE,
    SOURCE_FORMAT_DOCX,
    SOURCE_FORMAT_PDF,
    SOURCE_FORMAT_PPTX,
    STATUS_ACCEPTED,
    STATUS_PARTIAL,
    STATUS_UNRESOLVED,
    SourceSpan,
    StructuralEvidence,
)
from document_hierarchy_evidence_mapper import map_document_to_structural_evidence
from document_hierarchy_explicit_preserver import (
    preserve_explicit_docx_hierarchy,
    render_explicit_hierarchy,
)


def span(order: int) -> SourceSpan:
    position = SourcePosition("document_flow", 1, order, 0, 10)
    return SourceSpan(start=position, end=position)


def evidence(
    text: str,
    *,
    order: int,
    level: Optional[int] = None,
    role: str = EVIDENCE_ROLE_HEADING,
    content_kind: str = CONTENT_KIND_DOCUMENT_BOUNDARY,
    source_format: str = SOURCE_FORMAT_DOCX,
    numbering=None,
) -> StructuralEvidence:
    native = {"role_hint": ROLE_HEADING, "style_hint": f"Heading {level}", "native_hierarchy_hint": level}
    if level is None:
        native = {"role_hint": ROLE_BODY}
    return StructuralEvidence(
        evidence_id=f"ev-{order:04d}",
        source_format=source_format,
        source_order=order,
        source_span=span(order),
        raw_text=text,
        normalized_text=text,
        evidence_role=role,
        content_kind=content_kind,
        native_evidence=native,
        numbering_evidence=numbering or {},
        confidence_hint=CONFIDENCE_HIGH if level else CONFIDENCE_LOW,
    )


class ExplicitDocxHierarchyPreserverTests(unittest.TestCase):
    def test_simple_hierarchy(self):
        structure = preserve_explicit_docx_hierarchy(
            [
                evidence("A", order=1, level=1),
                evidence("B", order=2, level=2),
                evidence("C", order=3, level=2),
            ],
            document_id="doc-1",
            document_title="Simple",
        )

        self.assertEqual(structure.structure_status, STATUS_ACCEPTED)
        self.assertEqual(structure.structure_confidence, CONFIDENCE_HIGH)
        a, b, c = structure.nodes
        self.assertIsNone(a.parent_id)
        self.assertEqual(b.parent_id, a.node_id)
        self.assertEqual(c.parent_id, a.node_id)
        self.assertEqual(render_explicit_hierarchy(structure), "\n".join([
            "[accepted] L1 A",
            "  [accepted] L2 B",
            "  [accepted] L2 C",
        ]))

    def test_three_levels(self):
        structure = preserve_explicit_docx_hierarchy(
            [
                evidence("A", order=1, level=1),
                evidence("B", order=2, level=2),
                evidence("C", order=3, level=3),
                evidence("D", order=4, level=2),
            ],
            document_id="doc-2",
        )

        a, b, c, d = structure.nodes
        self.assertEqual(b.parent_id, a.node_id)
        self.assertEqual(c.parent_id, b.node_id)
        self.assertEqual(d.parent_id, a.node_id)

    def test_multiple_top_level_sections_reset_parent_stack(self):
        structure = preserve_explicit_docx_hierarchy(
            [
                evidence("A", order=1, level=1),
                evidence("B", order=2, level=2),
                evidence("C", order=3, level=1),
                evidence("D", order=4, level=2),
            ],
            document_id="doc-3",
        )

        a, b, c, d = structure.nodes
        self.assertEqual(b.parent_id, a.node_id)
        self.assertIsNone(c.parent_id)
        self.assertEqual(d.parent_id, c.node_id)

    def test_missing_parent_level_remains_unresolved_without_repair(self):
        structure = preserve_explicit_docx_hierarchy(
            [
                evidence("A", order=1, level=1),
                evidence("C", order=2, level=3),
            ],
            document_id="doc-4",
        )

        a, c = structure.nodes
        self.assertEqual(a.status, STATUS_ACCEPTED)
        self.assertEqual(c.level, 3)
        self.assertIsNone(c.parent_id)
        self.assertEqual(c.status, STATUS_UNRESOLVED)
        self.assertEqual(c.provenance.repair_reason, "missing_explicit_parent")
        self.assertEqual(structure.structure_status, STATUS_PARTIAL)
        self.assertEqual(structure.structure_confidence, CONFIDENCE_MEDIUM)
        self.assertEqual(structure.diagnostics["unresolved_parent_count"], 1)

    def test_document_starting_at_heading_2_preserves_level_without_promotion(self):
        structure = preserve_explicit_docx_hierarchy(
            [evidence("B", order=1, level=2)],
            document_id="doc-5",
        )

        node = structure.nodes[0]
        self.assertEqual(node.level, 2)
        self.assertIsNone(node.parent_id)
        self.assertEqual(node.status, STATUS_UNRESOLVED)

    def test_body_list_and_table_are_ignored(self):
        structure = preserve_explicit_docx_hierarchy(
            [
                evidence("A", order=1, level=1),
                evidence("Body", order=2, level=None, role=EVIDENCE_ROLE_BODY, content_kind=CONTENT_KIND_BODY),
                evidence(
                    "Local numbered list",
                    order=3,
                    level=0,
                    role=EVIDENCE_ROLE_LIST,
                    content_kind=CONTENT_KIND_LIST_ITEM,
                    numbering={"num_id": 5, "ilvl": 0},
                ),
                evidence("A | B", order=4, level=None, role=EVIDENCE_ROLE_TABLE, content_kind=CONTENT_KIND_TABLE),
                evidence("B", order=5, level=2),
            ],
            document_id="doc-6",
        )

        self.assertEqual([node.source_title for node in structure.nodes], ["A", "B"])
        self.assertEqual(structure.diagnostics["ignored_non_heading_count"], 3)

    def test_numbering_does_not_override_native_heading_style_level(self):
        structure = preserve_explicit_docx_hierarchy(
            [
                evidence("A", order=1, level=1, numbering={"num_id": 5, "ilvl": 0}),
                evidence("B", order=2, level=1, numbering={"num_id": 5, "ilvl": 1}),
                evidence("C", order=3, level=1, numbering={"num_id": 5, "ilvl": 2}),
            ],
            document_id="doc-7",
        )

        self.assertEqual([node.level for node in structure.nodes], [1, 1, 1])
        self.assertEqual([node.parent_id for node in structure.nodes], [None, None, None])
        self.assertEqual(structure.structure_status, STATUS_ACCEPTED)

    def test_provenance_preserves_source_span_and_supporting_evidence(self):
        source = evidence("A", order=1, level=1)
        structure = preserve_explicit_docx_hierarchy([source], document_id="doc-8")
        node = structure.nodes[0]

        self.assertEqual(node.source_span, source.source_span)
        self.assertEqual(node.evidence_ids, [source.evidence_id])
        self.assertEqual(node.provenance.supporting_evidence_ids, [source.evidence_id])
        self.assertEqual(node.provenance.decision_source, ORIGIN_EXPLICIT_SOURCE)
        self.assertEqual(node.origin, ORIGIN_EXPLICIT_SOURCE)
        self.assertEqual(node.confidence, CONFIDENCE_HIGH)

    def test_non_docx_evidence_is_ignored(self):
        structure = preserve_explicit_docx_hierarchy(
            [
                evidence("Slide title", order=1, level=1, source_format=SOURCE_FORMAT_PPTX),
                evidence("PDF large text", order=2, level=1, source_format=SOURCE_FORMAT_PDF),
            ],
            document_id="doc-9",
        )

        self.assertEqual(structure.nodes, [])
        self.assertEqual(structure.source_format, SOURCE_FORMAT_DOCX)
        self.assertEqual(structure.structure_status, STATUS_UNRESOLVED)
        self.assertEqual(structure.diagnostics["ignored_non_heading_count"], 2)

    def test_adapter_backed_docx_to_evidence_to_structure(self):
        try:
            from docx import Document
        except ImportError as exc:  # pragma: no cover - dependency check.
            raise unittest.SkipTest("python-docx is not installed") from exc

        document = Document()
        document.add_heading("Introduction", level=1)
        document.add_heading("Background", level=2)
        document.add_paragraph("Body is not promoted.")
        document.add_heading("Sources", level=3)
        buffer = io.BytesIO()
        document.save(buffer)

        canonical = adapt_docx_document(buffer.getvalue(), "headings.docx")
        mapped = map_document_to_structural_evidence(canonical)
        structure = preserve_explicit_docx_hierarchy(
            mapped,
            document_id="doc-10",
            document_title=canonical.title,
        )

        self.assertEqual([node.source_title for node in structure.nodes], ["Introduction", "Background", "Sources"])
        self.assertEqual([node.level for node in structure.nodes], [1, 2, 3])
        self.assertEqual(structure.nodes[1].parent_id, structure.nodes[0].node_id)
        self.assertEqual(structure.nodes[2].parent_id, structure.nodes[1].node_id)
        self.assertEqual(structure.diagnostics["ignored_non_heading_count"], 1)


if __name__ == "__main__":
    unittest.main()
