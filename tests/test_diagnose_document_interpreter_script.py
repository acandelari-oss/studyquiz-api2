import io
import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path


BACKEND_DIR = Path(__file__).resolve().parents[1]
SCRIPT = BACKEND_DIR / "scripts" / "diagnose_document_interpreter.py"


def _require_docx():
    try:
        from docx import Document
    except ImportError as exc:
        raise unittest.SkipTest("python-docx is not installed") from exc
    return Document


def _docx_bytes(document) -> bytes:
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


class DiagnoseDocumentInterpreterScriptTests(unittest.TestCase):
    def test_script_prints_readable_docx_diagnostic(self):
        Document = _require_docx()
        document = Document()
        document.add_heading("1 Cellular Biology", level=1)
        document.add_paragraph("Cells are the basic unit of life. " * 20)

        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "fixture.docx"
            path.write_bytes(_docx_bytes(document))
            completed = subprocess.run(
                [sys.executable, str(SCRIPT), str(path), "--max-chunks", "3"],
                cwd=str(BACKEND_DIR),
                text=True,
                capture_output=True,
                check=False,
            )

        self.assertEqual(completed.returncode, 0, completed.stderr)
        self.assertIn("DOUNO CLEAN DOCUMENT INTERPRETER DIAGNOSTIC", completed.stdout)
        self.assertIn("1 Cellular Biology", completed.stdout)
        self.assertIn("CHUNKS", completed.stdout)

    def test_script_prints_json_docx_diagnostic(self):
        Document = _require_docx()
        document = Document()
        document.add_heading("1 Gene Regulation", level=1)
        document.add_paragraph("Operons coordinate gene expression. " * 20)

        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "fixture.docx"
            path.write_bytes(_docx_bytes(document))
            completed = subprocess.run(
                [sys.executable, str(SCRIPT), str(path), "--json"],
                cwd=str(BACKEND_DIR),
                text=True,
                capture_output=True,
                check=False,
            )

        self.assertEqual(completed.returncode, 0, completed.stderr)
        payload = json.loads(completed.stdout)
        self.assertEqual(payload["document"]["source_format"], "docx")
        self.assertEqual(payload["sections"][0]["title"], "1 Gene Regulation")
        self.assertGreaterEqual(payload["summary"]["chunks"], 1)

    def test_missing_file_returns_clear_error(self):
        completed = subprocess.run(
            [sys.executable, str(SCRIPT), "/tmp/definitely-missing-douno-file.pdf"],
            cwd=str(BACKEND_DIR),
            text=True,
            capture_output=True,
            check=False,
        )

        self.assertEqual(completed.returncode, 2)
        self.assertIn("File not found", completed.stderr)


if __name__ == "__main__":
    unittest.main()
