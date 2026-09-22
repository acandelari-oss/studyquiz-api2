import io
import sys
import unittest
from pathlib import Path


BACKEND_DIR = Path(__file__).resolve().parents[1]
if str(BACKEND_DIR) not in sys.path:
    sys.path.insert(0, str(BACKEND_DIR))

from canonical_document import ROLE_BODY, ROLE_LIST, ROLE_TITLE, adapt_pptx_document
from document_hierarchy_evidence_mapper import map_document_to_structural_evidence


def _require_pptx():
    try:
        from pptx import Presentation
        from pptx.util import Inches
    except ImportError as exc:  # pragma: no cover - dependency check.
        raise unittest.SkipTest("python-pptx is not installed") from exc
    return Presentation, Inches


def _pptx_bytes(presentation) -> bytes:
    buffer = io.BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


def _canonical(presentation):
    return adapt_pptx_document(_pptx_bytes(presentation), "fixture.pptx")


class PptxLayoutEvidenceTests(unittest.TestCase):
    def test_standard_title_content_slide_preserves_layout_placeholder_and_geometry(self):
        Presentation, _Inches = _require_pptx()
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = "Cell metabolism"
        body = slide.placeholders[1]
        body.text = "ATP production"

        canonical = _canonical(presentation)
        evidence = map_document_to_structural_evidence(canonical)
        title = canonical.blocks[0]

        self.assertEqual(canonical.units[0].unit_type, "slide")
        self.assertEqual(canonical.units[0].unit_index, 1)
        self.assertEqual(canonical.units[0].display_label, "Cell metabolism")
        self.assertEqual(canonical.units[0].metadata["layout_name"], "Title and Content")
        self.assertEqual(title.role_hint, ROLE_TITLE)
        self.assertEqual(title.metadata["layout_name"], "Title and Content")
        self.assertEqual(title.metadata["layout_index"], 1)
        self.assertIn("TITLE", title.metadata["placeholder_type"])
        self.assertIn("placeholder_idx", title.metadata)
        for key in ("x", "y", "width", "height"):
            self.assertIsInstance(title.metadata[key], int)
            self.assertGreaterEqual(title.metadata[key], 0)

        self.assertEqual(evidence[0].native_evidence["layout_name"], "Title and Content")
        self.assertIn("TITLE", evidence[0].native_evidence["placeholder_type"])
        self.assertEqual(evidence[0].visual_evidence["x"], title.metadata["x"])
        self.assertEqual(evidence[0].visual_evidence["width"], title.metadata["width"])
        self.assertFalse(hasattr(evidence[0], "parent_id"))
        self.assertFalse(hasattr(evidence[0], "level"))

    def test_section_header_layout_is_preserved_without_section_detection(self):
        Presentation, _Inches = _require_pptx()
        presentation = Presentation()
        section_layout = next(
            (
                layout
                for layout in presentation.slide_layouts
                if "section" in (layout.name or "").lower()
            ),
            None,
        )
        if section_layout is None:
            self.skipTest("default template has no section-header-like layout")
        slide = presentation.slides.add_slide(section_layout)
        slide.shapes.title.text = "Gene regulation"

        canonical = _canonical(presentation)
        evidence = map_document_to_structural_evidence(canonical)

        self.assertEqual(canonical.blocks[0].metadata["layout_name"], section_layout.name)
        self.assertEqual(evidence[0].native_evidence["layout_name"], section_layout.name)
        self.assertEqual(evidence[0].evidence_role, "title")
        self.assertFalse(hasattr(evidence[0], "section"))
        self.assertFalse(hasattr(evidence[0], "parent_id"))

    def test_title_and_bullets_preserve_placeholder_bullet_and_geometry_evidence(self):
        Presentation, _Inches = _require_pptx()
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = "Respiration"
        text_frame = slide.placeholders[1].text_frame
        text_frame.clear()
        paragraph = text_frame.paragraphs[0]
        paragraph.text = "Aerobic respiration"
        paragraph.level = 1

        canonical = _canonical(presentation)
        title, bullet = canonical.blocks
        evidence = map_document_to_structural_evidence(canonical)

        self.assertEqual(title.role_hint, ROLE_TITLE)
        self.assertEqual(bullet.role_hint, ROLE_LIST)
        self.assertEqual(bullet.native_hierarchy_hint, 1)
        self.assertEqual(bullet.metadata["bullet_level"], 1)
        self.assertIn("placeholder_type", bullet.metadata)
        self.assertIn("placeholder_idx", bullet.metadata)
        self.assertIn("x", bullet.metadata)
        self.assertEqual(evidence[1].numbering_evidence["list_level"], 1)
        self.assertEqual(evidence[1].native_evidence["bullet_level"], 1)
        self.assertNotEqual(evidence[1].content_kind, "document_boundary_evidence")

    def test_nested_bullets_preserve_distinct_native_bullet_levels(self):
        Presentation, _Inches = _require_pptx()
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = "Nested bullets"
        text_frame = slide.placeholders[1].text_frame
        text_frame.clear()
        first = text_frame.paragraphs[0]
        first.text = "First nested item"
        first.level = 1
        second = text_frame.add_paragraph()
        second.text = "Second nested item"
        second.level = 2

        canonical = _canonical(presentation)
        bullet_blocks = [block for block in canonical.blocks if block.role_hint == ROLE_LIST]
        evidence = [
            item
            for item in map_document_to_structural_evidence(canonical)
            if item.evidence_role == "list"
        ]

        self.assertEqual([block.metadata["bullet_level"] for block in bullet_blocks], [1, 2])
        self.assertEqual([item.numbering_evidence["list_level"] for item in evidence], [1, 2])
        for item in evidence:
            self.assertFalse(hasattr(item, "parent_id"))

    def test_ordinary_text_box_preserves_geometry_without_placeholder_or_hierarchy(self):
        Presentation, Inches = _require_pptx()
        presentation = Presentation()
        blank_layout = presentation.slide_layouts[6]
        slide = presentation.slides.add_slide(blank_layout)
        box = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(3), Inches(1))
        box.text = "Side annotation"

        canonical = _canonical(presentation)
        evidence = map_document_to_structural_evidence(canonical)[0]
        block = canonical.blocks[0]

        self.assertEqual(block.role_hint, ROLE_BODY)
        self.assertNotIn("placeholder_type", block.metadata)
        self.assertEqual(block.metadata["layout_name"], "Blank")
        self.assertEqual(block.metadata["x"], int(Inches(1)))
        self.assertEqual(block.metadata["y"], int(Inches(2)))
        self.assertEqual(block.metadata["width"], int(Inches(3)))
        self.assertEqual(block.metadata["height"], int(Inches(1)))
        self.assertEqual(evidence.visual_evidence["x"], int(Inches(1)))
        self.assertFalse(hasattr(evidence, "parent_id"))

    def test_repeated_titles_remain_distinct_by_source_position(self):
        Presentation, _Inches = _require_pptx()
        presentation = Presentation()
        for _index in range(2):
            slide = presentation.slides.add_slide(presentation.slide_layouts[0])
            slide.shapes.title.text = "Action potential"

        canonical = _canonical(presentation)
        title_blocks = [block for block in canonical.blocks if block.role_hint == ROLE_TITLE]

        self.assertEqual([unit.display_label for unit in canonical.units], ["Action potential", "Action potential"])
        self.assertEqual([block.source_position.unit_index for block in title_blocks], [1, 2])
        self.assertEqual([block.source_order for block in title_blocks], [1, 2])

    def test_existing_title_behavior_uses_underlying_pptx_title_element(self):
        Presentation, _Inches = _require_pptx()
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[0])
        slide.shapes.title.text = "Cardiac Muscle"

        canonical = _canonical(presentation)

        self.assertEqual(canonical.units[0].display_label, "Cardiac Muscle")
        self.assertEqual(canonical.blocks[0].role_hint, ROLE_TITLE)


if __name__ == "__main__":
    unittest.main()
