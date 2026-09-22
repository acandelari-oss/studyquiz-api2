import io
import sys
import unittest
from pathlib import Path


BACKEND_DIR = Path(__file__).resolve().parents[1]
if str(BACKEND_DIR) not in sys.path:
    sys.path.insert(0, str(BACKEND_DIR))

from canonical_document import ROLE_BODY, ROLE_HEADING, ROLE_LIST, adapt_docx_document
from document_hierarchy_evidence_mapper import map_document_to_structural_evidence


def _require_docx():
    try:
        from docx import Document
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
    except ImportError as exc:  # pragma: no cover - dependency check.
        raise unittest.SkipTest("python-docx is not installed") from exc
    return Document, OxmlElement, qn


def _set_direct_numbering(paragraph, *, num_id: int = 5, ilvl: int = 0) -> None:
    _document, OxmlElement, qn = _require_docx()
    p_pr = paragraph._p.get_or_add_pPr()  # noqa: SLF001 - test fixture XML.
    num_pr = p_pr.numPr
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        p_pr.append(num_pr)

    ilvl_el = num_pr.find(qn("w:ilvl"))
    if ilvl_el is None:
        ilvl_el = OxmlElement("w:ilvl")
        num_pr.append(ilvl_el)
    ilvl_el.set(qn("w:val"), str(ilvl))

    num_id_el = num_pr.find(qn("w:numId"))
    if num_id_el is None:
        num_id_el = OxmlElement("w:numId")
        num_pr.append(num_id_el)
    num_id_el.set(qn("w:val"), str(num_id))


def _docx_bytes(document) -> bytes:
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


class DocxNumberingEvidenceTests(unittest.TestCase):
    def test_simple_numbered_list_preserves_native_numbering_evidence(self):
        Document, _OxmlElement, _qn = _require_docx()
        document = Document()
        paragraph = document.add_paragraph("Native list item without visible number")
        _set_direct_numbering(paragraph, num_id=5, ilvl=0)

        canonical = adapt_docx_document(_docx_bytes(document), "numbered.docx")
        block = canonical.blocks[0]
        evidence = map_document_to_structural_evidence(canonical)[0]

        self.assertEqual(block.role_hint, ROLE_LIST)
        self.assertEqual(block.native_hierarchy_hint, 0)
        self.assertEqual(block.metadata["numbering"]["source"], "direct")
        self.assertEqual(block.metadata["numbering"]["num_id"], 5)
        self.assertEqual(block.metadata["numbering"]["ilvl"], 0)
        self.assertEqual(block.metadata["numbering"]["abstract_num_id"], 7)
        self.assertEqual(block.metadata["numbering"]["format"], "decimal")
        self.assertEqual(evidence.numbering_evidence["source"], "direct")
        self.assertEqual(evidence.numbering_evidence["num_id"], 5)
        self.assertEqual(evidence.numbering_evidence["ilvl"], 0)
        self.assertEqual(evidence.numbering_evidence["list_level"], 0)
        self.assertFalse(hasattr(evidence, "parent_id"))
        self.assertFalse(hasattr(evidence, "level"))

    def test_nested_numbered_list_preserves_distinct_native_levels(self):
        Document, _OxmlElement, _qn = _require_docx()
        document = Document()
        parent = document.add_paragraph("Parent list item")
        child = document.add_paragraph("Child list item")
        _set_direct_numbering(parent, num_id=5, ilvl=0)
        _set_direct_numbering(child, num_id=5, ilvl=1)

        canonical = adapt_docx_document(_docx_bytes(document), "nested.docx")
        evidence = map_document_to_structural_evidence(canonical)

        self.assertEqual([block.metadata["numbering"]["ilvl"] for block in canonical.blocks], [0, 1])
        self.assertEqual([item.numbering_evidence["ilvl"] for item in evidence], [0, 1])
        self.assertEqual([item.numbering_evidence["list_level"] for item in evidence], [0, 1])
        for item in evidence:
            self.assertFalse(hasattr(item, "parent_id"))

    def test_numbered_heading_preserves_style_and_numbering_without_interpretation(self):
        Document, _OxmlElement, _qn = _require_docx()
        document = Document()
        heading = document.add_heading("Chapter title without visible number", level=1)
        _set_direct_numbering(heading, num_id=5, ilvl=0)

        canonical = adapt_docx_document(_docx_bytes(document), "numbered-heading.docx")
        block = canonical.blocks[0]
        evidence = map_document_to_structural_evidence(canonical)[0]

        self.assertEqual(block.role_hint, ROLE_HEADING)
        self.assertEqual(block.style_hint, "Heading 1")
        self.assertEqual(block.native_hierarchy_hint, 1)
        self.assertEqual(block.metadata["numbering"]["num_id"], 5)
        self.assertEqual(block.metadata["numbering"]["ilvl"], 0)
        self.assertEqual(evidence.native_evidence["style_hint"], "Heading 1")
        self.assertEqual(evidence.native_evidence["native_hierarchy_hint"], 1)
        self.assertEqual(evidence.numbering_evidence["num_id"], 5)
        self.assertEqual(evidence.numbering_evidence["ilvl"], 0)
        self.assertFalse(hasattr(evidence, "parent_id"))

    def test_ordinary_paragraph_remains_unnumbered(self):
        Document, _OxmlElement, _qn = _require_docx()
        document = Document()
        document.add_paragraph("Ordinary paragraph.")

        canonical = adapt_docx_document(_docx_bytes(document), "plain.docx")
        evidence = map_document_to_structural_evidence(canonical)[0]

        self.assertEqual(canonical.blocks[0].role_hint, ROLE_BODY)
        self.assertNotIn("numbering", canonical.blocks[0].metadata)
        self.assertEqual(evidence.numbering_evidence, {})

    def test_local_numbered_content_preserves_evidence_without_global_hierarchy(self):
        Document, _OxmlElement, _qn = _require_docx()
        document = Document()
        document.add_paragraph("Protocol")
        step = document.add_paragraph("Mix the solution")
        _set_direct_numbering(step, num_id=5, ilvl=0)

        canonical = adapt_docx_document(_docx_bytes(document), "local-list.docx")
        evidence = map_document_to_structural_evidence(canonical)

        self.assertEqual(canonical.blocks[1].role_hint, ROLE_LIST)
        self.assertEqual(evidence[1].numbering_evidence["num_id"], 5)
        self.assertEqual(evidence[1].numbering_evidence["ilvl"], 0)
        self.assertFalse(hasattr(evidence[1], "parent_id"))
        self.assertFalse(hasattr(evidence[1], "section"))

    def test_existing_heading_behavior_remains_unchanged(self):
        Document, _OxmlElement, _qn = _require_docx()
        document = Document()
        document.add_heading("Chapter", level=1)
        document.add_heading("Section", level=2)
        document.add_heading("Subsection", level=3)

        canonical = adapt_docx_document(_docx_bytes(document), "headings.docx")

        self.assertEqual(
            [(block.role_hint, block.style_hint, block.native_hierarchy_hint) for block in canonical.blocks],
            [
                (ROLE_HEADING, "Heading 1", 1),
                (ROLE_HEADING, "Heading 2", 2),
                (ROLE_HEADING, "Heading 3", 3),
            ],
        )
        self.assertEqual([block.source_order for block in canonical.blocks], [1, 2, 3])
        self.assertEqual([block.source_position.block_index for block in canonical.blocks], [1, 2, 3])


if __name__ == "__main__":
    unittest.main()
