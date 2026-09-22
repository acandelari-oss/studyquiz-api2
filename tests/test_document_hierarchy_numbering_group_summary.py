import io
import sys
import unittest
from pathlib import Path
from typing import Optional


BACKEND_DIR = Path(__file__).resolve().parents[1]
if str(BACKEND_DIR) not in sys.path:
    sys.path.insert(0, str(BACKEND_DIR))

from canonical_document import SourcePosition, adapt_docx_document
from document_hierarchy_consistency import analyze_docx_heading_numbering_consistency
from document_hierarchy_contracts import (
    CONFIDENCE_HIGH,
    CONTENT_KIND_BODY,
    CONTENT_KIND_DOCUMENT_BOUNDARY,
    CONTENT_KIND_LIST_ITEM,
    CONTENT_KIND_TABLE,
    EVIDENCE_ROLE_BODY,
    EVIDENCE_ROLE_HEADING,
    EVIDENCE_ROLE_LIST,
    EVIDENCE_ROLE_TABLE,
    SOURCE_FORMAT_DOCX,
    SourceSpan,
    StructuralEvidence,
)
from document_hierarchy_evidence_mapper import map_document_to_structural_evidence
from document_hierarchy_explicit_preserver import preserve_explicit_docx_hierarchy
from document_hierarchy_numbering_authority import (
    AUTHORITY_TRUSTED_NUMBERING_SYSTEM,
    classify_docx_numbering_authority,
)
from document_hierarchy_numbering_group_summary import (
    render_docx_numbering_group_summary,
    summarize_docx_numbering_groups,
)
from document_hierarchy_numbering_observer import observe_docx_numbering_patterns


def span(order: int, *, unit_index: int = 1) -> SourceSpan:
    position = SourcePosition("document_flow", unit_index, order, 0, 10)
    return SourceSpan(start=position, end=position)


def heading(
    order: int,
    *,
    level: int,
    ilvl: Optional[int],
    num_id=5,
    abstract_num_id=7,
    unit_index: int = 1,
) -> StructuralEvidence:
    numbering = {}
    if ilvl is not None or num_id is not None or abstract_num_id is not None:
        numbering = {
            "source": "direct",
            "format": "decimal",
        }
        if num_id is not None:
            numbering["num_id"] = num_id
        if abstract_num_id is not None:
            numbering["abstract_num_id"] = abstract_num_id
        if ilvl is not None:
            numbering["ilvl"] = ilvl
            numbering["list_level"] = ilvl
    return StructuralEvidence(
        evidence_id=f"ev-{order:04d}",
        source_format=SOURCE_FORMAT_DOCX,
        source_order=order,
        source_span=span(order, unit_index=unit_index),
        raw_text=f"Heading {order}",
        normalized_text=f"Heading {order}",
        evidence_role=EVIDENCE_ROLE_HEADING,
        content_kind=CONTENT_KIND_DOCUMENT_BOUNDARY,
        native_evidence={
            "role_hint": "heading",
            "style_hint": f"Heading {level}",
            "native_hierarchy_hint": level,
        },
        numbering_evidence=numbering,
        structural_context={
            "unit_type": "document_flow",
            "unit_index": unit_index,
            "block_index": order,
        },
        confidence_hint=CONFIDENCE_HIGH,
    )


def numbered(
    order: int,
    *,
    ilvl: Optional[int],
    num_id=5,
    abstract_num_id=7,
    source="direct",
    start_override=None,
    evidence_role=EVIDENCE_ROLE_LIST,
    content_kind=CONTENT_KIND_LIST_ITEM,
    unit_index: int = 1,
) -> StructuralEvidence:
    numbering = {
        "source": source,
        "format": "decimal",
    }
    if num_id is not None:
        numbering["num_id"] = num_id
    if abstract_num_id is not None:
        numbering["abstract_num_id"] = abstract_num_id
    if ilvl is not None:
        numbering["ilvl"] = ilvl
        numbering["list_level"] = ilvl
    if start_override is not None:
        numbering["start_override"] = start_override
    return StructuralEvidence(
        evidence_id=f"ev-{order:04d}",
        source_format=SOURCE_FORMAT_DOCX,
        source_order=order,
        source_span=span(order, unit_index=unit_index),
        raw_text=f"Numbered {order}",
        normalized_text=f"Numbered {order}",
        evidence_role=evidence_role,
        content_kind=content_kind,
        native_evidence={"role_hint": "list", "native_hierarchy_hint": ilvl},
        numbering_evidence=numbering,
        structural_context={
            "unit_type": "document_flow",
            "unit_index": unit_index,
            "block_index": order,
        },
        confidence_hint=CONFIDENCE_HIGH,
    )


def body(order: int) -> StructuralEvidence:
    return StructuralEvidence(
        evidence_id=f"ev-{order:04d}",
        source_format=SOURCE_FORMAT_DOCX,
        source_order=order,
        source_span=span(order),
        raw_text=f"Body {order}",
        normalized_text=f"Body {order}",
        evidence_role=EVIDENCE_ROLE_BODY,
        content_kind=CONTENT_KIND_BODY,
        native_evidence={"role_hint": "body"},
    )


def summarize(items):
    structure = preserve_explicit_docx_hierarchy(items, document_id="doc")
    consistency = analyze_docx_heading_numbering_consistency(items, structure)
    observations = observe_docx_numbering_patterns(items)
    authority = classify_docx_numbering_authority(items, structure, consistency, observations)
    summary = summarize_docx_numbering_groups(
        items,
        structure,
        consistency,
        observations,
        authority,
    )
    return structure, consistency, observations, authority, summary


def _require_docx():
    try:
        from docx import Document
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
    except ImportError as exc:  # pragma: no cover - dependency check.
        raise unittest.SkipTest("python-docx is not installed") from exc
    return Document, OxmlElement, qn


def _set_direct_numbering(paragraph, *, num_id: int = 5, ilvl: int = 0) -> None:
    _Document, OxmlElement, qn = _require_docx()
    p_pr = paragraph._p.get_or_add_pPr()  # noqa: SLF001 - test fixture XML.
    num_pr = p_pr.numPr
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        p_pr.append(num_pr)
    ilvl_el = num_pr.find(qn("w:ilvl"))
    if ilvl_el is None:
        ilvl_el = OxmlElement("w:ilvl")
        num_pr.append(ilvl_el)
    ilvl_el.set(qn("w:val"), str(ilvl))
    num_id_el = num_pr.find(qn("w:numId"))
    if num_id_el is None:
        num_id_el = OxmlElement("w:numId")
        num_pr.append(num_id_el)
    num_id_el.set(qn("w:val"), str(num_id))


class DocxNumberingGroupSummaryTests(unittest.TestCase):
    def test_heading_only_group_summarizes_confirmed_pairs_without_trust(self):
        _structure, _consistency, _observations, _authority, summary = summarize([
            heading(1, level=1, ilvl=0, num_id=5),
            heading(2, level=2, ilvl=1, num_id=5),
            heading(3, level=3, ilvl=2, num_id=5),
        ])

        group = summary.groups[0]
        self.assertEqual(group.num_id, 5)
        self.assertEqual(group.total_evidence_count, 3)
        self.assertEqual(group.explicit_heading_count, 3)
        self.assertEqual(group.non_heading_count, 0)
        self.assertEqual(group.explicit_confirmed_count, 3)
        self.assertEqual(group.heading_depth_pairs, {"H1:D1": 1, "H2:D2": 1, "H3:D3": 1})
        self.assertEqual(group.trusted_numbering_system_count, 0)

    def test_mixed_confirmed_and_conflicting_headings_are_exposed(self):
        _structure, _consistency, _observations, _authority, summary = summarize([
            heading(1, level=1, ilvl=0, num_id=5),
            heading(2, level=1, ilvl=1, num_id=5),
        ])

        group = summary.groups[0]
        self.assertEqual(group.explicit_confirmed_count, 1)
        self.assertEqual(group.explicit_conflict_count, 1)
        self.assertEqual(group.heading_depth_pairs, {"H1:D1": 1, "H1:D2": 1})
        self.assertEqual(summary.diagnostics["groups_with_conflicts"], 1)

    def test_non_heading_only_group_summarizes_ambiguous_observations_and_no_nodes(self):
        structure, _consistency, _observations, _authority, summary = summarize([
            numbered(1, ilvl=0, num_id=5),
            numbered(2, ilvl=1, num_id=5),
        ])

        group = summary.groups[0]
        self.assertEqual(structure.nodes, [])
        self.assertEqual(group.explicit_heading_count, 0)
        self.assertEqual(group.non_heading_count, 2)
        self.assertEqual(group.ambiguous_count, 2)
        self.assertEqual(group.observed_depths, [1, 2])

    def test_mixed_heading_and_non_heading_group_keeps_both_populations_visible(self):
        _structure, _consistency, _observations, _authority, summary = summarize([
            heading(1, level=1, ilvl=0, num_id=5),
            numbered(2, ilvl=1, num_id=5),
        ])

        group = summary.groups[0]
        self.assertEqual(group.explicit_heading_count, 1)
        self.assertEqual(group.non_heading_count, 1)
        self.assertEqual(group.explicit_confirmed_count, 1)
        self.assertEqual(group.ambiguous_count, 1)
        self.assertEqual(summary.diagnostics["mixed_groups"], 1)

    def test_local_table_evidence_is_not_hidden_by_confirmed_headings(self):
        _structure, _consistency, _observations, _authority, summary = summarize([
            heading(1, level=1, ilvl=0, num_id=5),
            numbered(
                2,
                ilvl=0,
                num_id=5,
                evidence_role=EVIDENCE_ROLE_TABLE,
                content_kind=CONTENT_KIND_TABLE,
            ),
        ])

        group = summary.groups[0]
        self.assertEqual(group.explicit_confirmed_count, 1)
        self.assertEqual(group.local_or_unsupported_count, 1)
        self.assertEqual(summary.diagnostics["groups_with_local_or_unsupported_evidence"], 1)

    def test_separate_num_id_values_sharing_abstract_num_id_remain_separate(self):
        _structure, _consistency, _observations, _authority, summary = summarize([
            numbered(1, ilvl=0, num_id=5, abstract_num_id=2),
            numbered(2, ilvl=0, num_id=8, abstract_num_id=2),
        ])

        self.assertEqual([group.num_id for group in summary.groups], [5, 8])
        self.assertEqual([group.abstract_num_ids for group in summary.groups], [[2], [2]])
        self.assertEqual(summary.diagnostics["numbering_group_count"], 2)

    def test_multiple_depths_and_transitions_are_aggregated(self):
        _structure, _consistency, _observations, _authority, summary = summarize([
            numbered(1, ilvl=0, num_id=5),
            numbered(2, ilvl=1, num_id=5),
            numbered(3, ilvl=2, num_id=5),
            numbered(4, ilvl=1, num_id=5),
        ])

        group = summary.groups[0]
        self.assertEqual(group.observed_ilvls, [0, 1, 2])
        self.assertEqual(group.observed_depths, [1, 2, 3])
        self.assertEqual(group.transition_types, ["ascend", "descend", "start"])

    def test_interruption_is_aggregated(self):
        _structure, _consistency, _observations, _authority, summary = summarize([
            numbered(1, ilvl=0, num_id=5),
            body(2),
            numbered(3, ilvl=0, num_id=5),
        ])

        self.assertEqual(summary.groups[0].interruption_count, 1)

    def test_restart_is_aggregated(self):
        _structure, _consistency, _observations, _authority, summary = summarize([
            numbered(1, ilvl=0, num_id=5),
            numbered(2, ilvl=0, num_id=5, start_override=1),
        ])

        group = summary.groups[0]
        self.assertEqual(group.restart_count, 1)
        self.assertIn("restart", group.transition_types)

    def test_source_order_span_and_units_are_factual(self):
        _structure, _consistency, _observations, _authority, summary = summarize([
            numbered(7, ilvl=0, num_id=5, unit_index=1),
            numbered(20, ilvl=0, num_id=5, unit_index=3),
        ])

        group = summary.groups[0]
        self.assertEqual(group.min_source_order, 7)
        self.assertEqual(group.max_source_order, 20)
        self.assertEqual(group.distinct_unit_count, 2)
        self.assertEqual(group.first_unit, "document_flow:1")
        self.assertEqual(group.last_unit, "document_flow:3")

    def test_ungroupable_evidence_missing_num_id_is_not_merged(self):
        _structure, _consistency, _observations, _authority, summary = summarize([
            numbered(1, ilvl=0, num_id=None, abstract_num_id=7),
            numbered(2, ilvl=0, num_id=5, abstract_num_id=7),
        ])

        self.assertEqual([group.num_id for group in summary.groups], [5])
        self.assertEqual(len(summary.ungroupable_evidence), 1)
        self.assertEqual(summary.ungroupable_evidence[0].reason, "missing_num_id")
        self.assertEqual(summary.ungroupable_evidence[0].abstract_num_id, 7)

    def test_provenance_evidence_and_node_ids_survive(self):
        _structure, _consistency, _observations, _authority, summary = summarize([
            heading(1, level=1, ilvl=0, num_id=5),
            numbered(2, ilvl=1, num_id=5),
        ])

        group = summary.groups[0]
        self.assertEqual(group.evidence_ids, ["ev-0001", "ev-0002"])
        self.assertEqual(group.node_ids, ["docx-heading-0001"])

    def test_step_9_structure_is_immutable(self):
        items = [heading(1, level=1, ilvl=0, num_id=5), numbered(2, ilvl=1, num_id=5)]
        before = preserve_explicit_docx_hierarchy(items, document_id="doc")
        summarize(items)
        after = preserve_explicit_docx_hierarchy(items, document_id="doc")

        self.assertEqual(before, after)

    def test_step_10_11_13_reports_are_immutable(self):
        items = [heading(1, level=1, ilvl=0, num_id=5), numbered(2, ilvl=1, num_id=5)]
        structure = preserve_explicit_docx_hierarchy(items, document_id="doc")
        consistency_before = analyze_docx_heading_numbering_consistency(items, structure)
        observations_before = observe_docx_numbering_patterns(items)
        authority_before = classify_docx_numbering_authority(
            items,
            structure,
            consistency_before,
            observations_before,
        )
        summarize_docx_numbering_groups(
            items,
            structure,
            consistency_before,
            observations_before,
            authority_before,
        )
        consistency_after = analyze_docx_heading_numbering_consistency(items, structure)
        observations_after = observe_docx_numbering_patterns(items)
        authority_after = classify_docx_numbering_authority(
            items,
            structure,
            consistency_after,
            observations_after,
        )

        self.assertEqual(consistency_before, consistency_after)
        self.assertEqual(observations_before, observations_after)
        self.assertEqual(authority_before, authority_after)

    def test_trusted_numbering_system_remains_unassigned_for_strong_looking_group(self):
        _structure, _consistency, _observations, authority, summary = summarize([
            heading(1, level=1, ilvl=0, num_id=5),
            heading(2, level=2, ilvl=1, num_id=5),
            heading(3, level=3, ilvl=2, num_id=5),
            heading(4, level=1, ilvl=0, num_id=5),
            heading(5, level=2, ilvl=1, num_id=5),
            numbered(6, ilvl=2, num_id=5),
        ])

        self.assertNotIn(
            AUTHORITY_TRUSTED_NUMBERING_SYSTEM,
            {record.authority_class for record in authority.records},
        )
        self.assertEqual(summary.groups[0].trusted_numbering_system_count, 0)

    def test_adapter_backed_docx_flow_produces_group_summary(self):
        Document, _OxmlElement, _qn = _require_docx()
        document = Document()
        heading_one = document.add_heading("Introduction", level=1)
        heading_two = document.add_heading("Background", level=2)
        list_item = document.add_paragraph("Nested numbered body item")
        _set_direct_numbering(heading_one, num_id=5, ilvl=0)
        _set_direct_numbering(heading_two, num_id=5, ilvl=1)
        _set_direct_numbering(list_item, num_id=5, ilvl=2)
        buffer = io.BytesIO()
        document.save(buffer)

        canonical = adapt_docx_document(buffer.getvalue(), "summary.docx")
        evidence = map_document_to_structural_evidence(canonical)
        structure = preserve_explicit_docx_hierarchy(evidence, document_id="doc")
        consistency = analyze_docx_heading_numbering_consistency(evidence, structure)
        observations = observe_docx_numbering_patterns(evidence)
        authority = classify_docx_numbering_authority(evidence, structure, consistency, observations)
        summary = summarize_docx_numbering_groups(
            evidence,
            structure,
            consistency,
            observations,
            authority,
        )

        group = summary.groups[0]
        self.assertEqual(group.explicit_confirmed_count, 2)
        self.assertEqual(group.ambiguous_count, 1)
        self.assertEqual(group.observed_depths, [1, 2, 3])
        self.assertEqual([node.level for node in structure.nodes], [1, 2])

    def test_renderer_is_factual_and_does_not_add_trust(self):
        _structure, _consistency, _observations, _authority, summary = summarize([
            heading(1, level=1, ilvl=0, num_id=5),
            numbered(2, ilvl=1, num_id=5),
        ])

        rendered = render_docx_numbering_group_summary(summary)
        self.assertIn("NUMBERING GROUP num_id=5", rendered)
        self.assertIn("confirmed=1", rendered)
        self.assertIn("ambiguous=1", rendered)
        self.assertNotIn("trusted=True", rendered)
        self.assertNotIn("confidence", rendered)


if __name__ == "__main__":
    unittest.main()
