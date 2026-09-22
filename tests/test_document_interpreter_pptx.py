import io
import sys
import unittest
from pathlib import Path


BACKEND_DIR = Path(__file__).resolve().parents[1]
if str(BACKEND_DIR) not in sys.path:
    sys.path.insert(0, str(BACKEND_DIR))

from document_hierarchy_contracts import CONFIDENCE_HIGH, CONFIDENCE_MEDIUM
from document_interpreter_contract import (
    CONTENT_ROLE_BODY,
    CONTENT_ROLE_HEADING,
    CONTENT_ROLE_LIST,
    CONTENT_ROLE_TABLE,
    INTERPRETER_STATUS_ACCEPTED,
    INTERPRETER_STATUS_PARTIAL,
)
from document_interpreter_pptx import interpret_pptx_document


def _require_pptx():
    try:
        from pptx import Presentation
        from pptx.util import Inches
    except ImportError as exc:
        raise unittest.SkipTest("python-pptx is not installed") from exc
    return Presentation, Inches


def _pptx_bytes(presentation) -> bytes:
    buffer = io.BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


class DocumentInterpreterPptxTests(unittest.TestCase):
    def test_title_and_content_slide_becomes_slide_owned_section(self):
        Presentation, _Inches = _require_pptx()
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = "Cell metabolism"
        slide.placeholders[1].text = "ATP production happens in mitochondria. " * 8

        result = interpret_pptx_document(
            _pptx_bytes(presentation),
            "slides.pptx",
            document_id="doc-1",
        )

        self.assertEqual(result.status, INTERPRETER_STATUS_ACCEPTED)
        self.assertEqual(result.confidence, CONFIDENCE_HIGH)
        self.assertEqual(len(result.sections), 1)
        section = result.sections[0]
        self.assertEqual(section.section_id, "pptx-slide-0001")
        self.assertEqual(section.title, "Cell metabolism")
        self.assertEqual(section.level, 1)
        self.assertEqual(section.confidence, CONFIDENCE_HIGH)
        self.assertEqual([block.section_id for block in section.owned_blocks], [section.section_id] * 2)
        self.assertEqual([block.role for block in section.owned_blocks], [CONTENT_ROLE_HEADING, CONTENT_ROLE_BODY])
        self.assertEqual([block.page for block in section.owned_blocks], [1, 1])
        self.assertTrue(result.diagnostics["contract_validation"]["is_valid"])

    def test_each_slide_remains_a_distinct_source_section(self):
        Presentation, _Inches = _require_pptx()
        presentation = Presentation()
        first = presentation.slides.add_slide(presentation.slide_layouts[1])
        first.shapes.title.text = "Respiration"
        first.placeholders[1].text = "Respiration body " * 12
        second = presentation.slides.add_slide(presentation.slide_layouts[1])
        second.shapes.title.text = "Photosynthesis"
        second.placeholders[1].text = "Photosynthesis body " * 12

        result = interpret_pptx_document(
            _pptx_bytes(presentation),
            "deck.pptx",
            document_id="doc-1",
        )

        self.assertEqual([section.title for section in result.sections], ["Respiration", "Photosynthesis"])
        self.assertEqual([section.section_id for section in result.sections], ["pptx-slide-0001", "pptx-slide-0002"])
        self.assertLess(result.sections[0].source_order, result.sections[1].source_order)

    def test_bullets_and_tables_are_owned_by_current_slide_section(self):
        Presentation, Inches = _require_pptx()
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = "Transport"
        frame = slide.placeholders[1].text_frame
        frame.clear()
        first = frame.paragraphs[0]
        first.text = "Passive transport"
        first.level = 1
        table_shape = slide.shapes.add_table(
            2,
            2,
            Inches(1),
            Inches(4),
            Inches(4),
            Inches(1),
        )
        table = table_shape.table
        table.cell(0, 0).text = "Type"
        table.cell(0, 1).text = "Energy"
        table.cell(1, 0).text = "Active"
        table.cell(1, 1).text = "ATP"

        result = interpret_pptx_document(
            _pptx_bytes(presentation),
            "table.pptx",
            document_id="doc-1",
        )

        roles = [block.role for block in result.sections[0].owned_blocks]
        self.assertIn(CONTENT_ROLE_LIST, roles)
        self.assertIn(CONTENT_ROLE_TABLE, roles)
        self.assertEqual({block.section_id for block in result.sections[0].owned_blocks}, {"pptx-slide-0001"})

    def test_titleless_slide_preserves_boundary_but_marks_partial(self):
        Presentation, Inches = _require_pptx()
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
        box.text = "This slide has body content but no title. " * 8

        result = interpret_pptx_document(
            _pptx_bytes(presentation),
            "titleless.pptx",
            document_id="doc-1",
        )

        self.assertEqual(result.status, INTERPRETER_STATUS_PARTIAL)
        self.assertEqual(result.confidence, CONFIDENCE_MEDIUM)
        self.assertEqual(result.sections[0].title, "Slide 1")
        self.assertEqual(result.sections[0].confidence, CONFIDENCE_MEDIUM)
        self.assertFalse(result.sections[0].metadata["has_explicit_title"])
        self.assertEqual(result.diagnostics["titleless_slide_count"], 1)
        self.assertTrue(result.diagnostics["contract_validation"]["is_valid"])


if __name__ == "__main__":
    unittest.main()
