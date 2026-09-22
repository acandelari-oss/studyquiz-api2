import io
import sys
import unittest
from pathlib import Path


BACKEND_DIR = Path(__file__).resolve().parents[1]
if str(BACKEND_DIR) not in sys.path:
    sys.path.insert(0, str(BACKEND_DIR))

from document_hierarchy_contracts import CONFIDENCE_HIGH, CONFIDENCE_LOW, CONFIDENCE_MEDIUM
from document_interpreter_contract import (
    CONTENT_ROLE_BODY,
    CONTENT_ROLE_HEADING,
    CONTENT_ROLE_LIST,
    CONTENT_ROLE_TABLE,
    INTERPRETER_STATUS_ACCEPTED,
    INTERPRETER_STATUS_FALLBACK,
    INTERPRETER_STATUS_PARTIAL,
)
from document_interpreter_docx import interpret_docx_document


def _require_docx():
    try:
        from docx import Document
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
    except ImportError as exc:  # pragma: no cover - dependency check.
        raise unittest.SkipTest("python-docx is not installed") from exc
    return Document, OxmlElement, qn


def _set_direct_numbering(paragraph, *, num_id: int = 5, ilvl: int = 0) -> None:
    _Document, OxmlElement, qn = _require_docx()
    p_pr = paragraph._p.get_or_add_pPr()  # noqa: SLF001 - test fixture XML.
    num_pr = p_pr.numPr
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        p_pr.append(num_pr)

    ilvl_el = num_pr.find(qn("w:ilvl"))
    if ilvl_el is None:
        ilvl_el = OxmlElement("w:ilvl")
        num_pr.append(ilvl_el)
    ilvl_el.set(qn("w:val"), str(ilvl))

    num_id_el = num_pr.find(qn("w:numId"))
    if num_id_el is None:
        num_id_el = OxmlElement("w:numId")
        num_pr.append(num_id_el)
    num_id_el.set(qn("w:val"), str(num_id))


def _docx_bytes(document) -> bytes:
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


class DocxDocumentInterpreterTests(unittest.TestCase):
    def test_heading_hierarchy_owns_following_content(self):
        Document, _OxmlElement, _qn = _require_docx()
        document = Document()
        document.add_heading("Chapter 1", level=1)
        document.add_paragraph("Chapter body.")
        document.add_heading("Section 1.1", level=2)
        document.add_paragraph("Section body.")

        result = interpret_docx_document(
            _docx_bytes(document),
            "sample.docx",
            document_id="doc-1",
        )

        self.assertEqual(result.status, INTERPRETER_STATUS_ACCEPTED)
        self.assertEqual(result.confidence, CONFIDENCE_HIGH)
        self.assertTrue(result.has_authoritative_structure)
        self.assertEqual([section.title for section in result.sections], ["Chapter 1", "Section 1.1"])
        self.assertIsNone(result.sections[0].parent_id)
        self.assertEqual(result.sections[1].parent_id, result.sections[0].section_id)
        self.assertEqual(
            [block.text for block in result.sections[0].owned_blocks],
            ["Chapter 1", "Chapter body."],
        )
        self.assertEqual(
            [block.text for block in result.sections[1].owned_blocks],
            ["Section 1.1", "Section body."],
        )
        self.assertEqual(result.diagnostics["contract_validation"]["is_valid"], True)

    def test_preface_before_first_heading_remains_unresolved(self):
        Document, _OxmlElement, _qn = _require_docx()
        document = Document()
        document.add_paragraph("Course preface.")
        document.add_heading("Main topic", level=1)
        document.add_paragraph("Main body.")

        result = interpret_docx_document(
            _docx_bytes(document),
            "preface.docx",
            document_id="doc-2",
        )

        self.assertEqual(result.status, INTERPRETER_STATUS_PARTIAL)
        self.assertEqual(result.confidence, CONFIDENCE_MEDIUM)
        self.assertEqual([block.text for block in result.unresolved_blocks], ["Course preface."])
        self.assertEqual(result.unresolved_blocks[0].confidence, CONFIDENCE_LOW)
        self.assertEqual(result.sections[0].owned_blocks[0].text, "Main topic")

    def test_no_heading_styles_returns_fallback_with_all_content(self):
        Document, _OxmlElement, _qn = _require_docx()
        document = Document()
        document.add_paragraph("Plain introduction.")
        document.add_paragraph("Plain body.")

        result = interpret_docx_document(
            _docx_bytes(document),
            "plain.docx",
            document_id="doc-3",
        )

        self.assertEqual(result.status, INTERPRETER_STATUS_FALLBACK)
        self.assertEqual(result.confidence, CONFIDENCE_LOW)
        self.assertEqual(result.sections, [])
        self.assertEqual([block.text for block in result.unresolved_blocks], ["Plain introduction.", "Plain body."])
        self.assertEqual(result.diagnostics["fallback_reason"], "docx_contains_no_heading_styles")

    def test_lists_and_tables_are_owned_by_current_section(self):
        Document, _OxmlElement, _qn = _require_docx()
        document = Document()
        document.add_heading("Methods", level=1)
        list_item = document.add_paragraph("First item")
        _set_direct_numbering(list_item)
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "A"
        table.cell(0, 1).text = "B"

        result = interpret_docx_document(
            _docx_bytes(document),
            "table.docx",
            document_id="doc-4",
        )

        self.assertEqual(result.status, INTERPRETER_STATUS_ACCEPTED)
        roles = [block.role for block in result.sections[0].owned_blocks]
        self.assertEqual(roles[0], CONTENT_ROLE_HEADING)
        self.assertIn(CONTENT_ROLE_LIST, roles)
        self.assertIn(CONTENT_ROLE_TABLE, roles)
        self.assertEqual([block.section_id for block in result.sections[0].owned_blocks], [result.sections[0].section_id] * 3)

    def test_heading_level_gap_is_partial_but_content_is_preserved(self):
        Document, _OxmlElement, _qn = _require_docx()
        document = Document()
        document.add_heading("Deep section", level=3)
        document.add_paragraph("Deep body.")

        result = interpret_docx_document(
            _docx_bytes(document),
            "gap.docx",
            document_id="doc-5",
        )

        self.assertEqual(result.status, INTERPRETER_STATUS_PARTIAL)
        self.assertEqual(result.confidence, CONFIDENCE_MEDIUM)
        self.assertEqual(result.sections[0].level, 3)
        self.assertIsNone(result.sections[0].parent_id)
        self.assertEqual(result.sections[0].confidence, CONFIDENCE_MEDIUM)
        self.assertEqual(result.diagnostics["heading_level_gaps"], 1)
        self.assertEqual([block.text for block in result.all_content_blocks], ["Deep section", "Deep body."])

    def test_multiple_heading_one_sections_are_distinct_top_level_sections(self):
        Document, _OxmlElement, _qn = _require_docx()
        document = Document()
        document.add_heading("Chapter A", level=1)
        document.add_paragraph("A body.")
        document.add_heading("Chapter B", level=1)
        document.add_paragraph("B body.")

        result = interpret_docx_document(
            _docx_bytes(document),
            "chapters.docx",
            document_id="doc-6",
        )

        self.assertEqual(result.status, INTERPRETER_STATUS_ACCEPTED)
        self.assertEqual([section.parent_id for section in result.sections], [None, None])
        self.assertEqual(
            [[block.text for block in section.owned_blocks] for section in result.sections],
            [["Chapter A", "A body."], ["Chapter B", "B body."]],
        )

    def test_content_roles_include_body_and_heading(self):
        Document, _OxmlElement, _qn = _require_docx()
        document = Document()
        document.add_heading("Topic", level=1)
        document.add_paragraph("Body.")

        result = interpret_docx_document(
            _docx_bytes(document),
            "roles.docx",
            document_id="doc-7",
        )

        self.assertEqual(
            [block.role for block in result.sections[0].owned_blocks],
            [CONTENT_ROLE_HEADING, CONTENT_ROLE_BODY],
        )


if __name__ == "__main__":
    unittest.main()
