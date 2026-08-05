import io
import os
import re
from dataclasses import dataclass
from typing import List, Optional

from pypdf import PdfReader


class DocumentExtractionError(ValueError):
    pass


@dataclass
class ExtractedBlock:
    text: str
    raw_text: str
    page: Optional[int]
    section: Optional[str] = None
    block_index: int = 0


@dataclass
class ExtractedDocument:
    filename: str
    file_format: str
    file_size_bytes: int
    blocks: List[ExtractedBlock]
    pages_detected: Optional[int] = None


def _normalize_extension(filename: str) -> str:
    return os.path.splitext(filename or "")[1].lower().strip()


def extract_uploaded_document(file_bytes: bytes, filename: str) -> ExtractedDocument:
    extension = _normalize_extension(filename)

    if extension == ".pdf":
        return extract_pdf_document(file_bytes, filename)

    if extension == ".docx":
        return extract_docx_document(file_bytes, filename)

    if extension == ".pptx":
        return extract_pptx_document(file_bytes, filename)

    raise DocumentExtractionError(
        "Unsupported document format. Please upload a PDF, DOCX, or PPTX file."
    )


def extract_pdf_document(file_bytes: bytes, filename: str) -> ExtractedDocument:
    try:
        pdf_stream = io.BytesIO(file_bytes)
        reader = PdfReader(pdf_stream)
        pages_detected = len(reader.pages)
    except Exception as exc:
        raise DocumentExtractionError(
            "The uploaded PDF could not be opened. The file may be corrupted."
        ) from exc

    if pages_detected <= 0:
        raise DocumentExtractionError("The uploaded PDF does not contain any pages.")

    blocks = []
    for page_index, page in enumerate(reader.pages):
        raw_page_text = page.extract_text()
        blocks.append(
            ExtractedBlock(
                text=raw_page_text or "",
                raw_text=raw_page_text or "",
                page=page_index + 1,
                block_index=page_index + 1,
            )
        )

    return ExtractedDocument(
        filename=filename,
        file_format="pdf",
        file_size_bytes=len(file_bytes),
        blocks=blocks,
        pages_detected=pages_detected,
    )


def extract_docx_document(file_bytes: bytes, filename: str) -> ExtractedDocument:
    try:
        from docx import Document
        from docx.table import Table
        from docx.text.paragraph import Paragraph
    except ImportError as exc:
        raise DocumentExtractionError(
            "DOCX support is not installed on this server."
        ) from exc

    try:
        document = Document(io.BytesIO(file_bytes))
    except Exception as exc:
        raise DocumentExtractionError(
            "The uploaded DOCX could not be opened. The file may be corrupted."
        ) from exc

    def iter_body_items(doc):
        for child in doc.element.body.iterchildren():
            if child.tag.endswith("}p"):
                yield Paragraph(child, doc)
            elif child.tag.endswith("}tbl"):
                yield Table(child, doc)

    def paragraph_text(paragraph) -> str:
        text = (paragraph.text or "").strip()
        if not text:
            return ""

        style_name = (getattr(paragraph.style, "name", "") or "").strip()
        if style_name.lower().startswith("heading"):
            return text

        return text

    def heading_level(paragraph) -> Optional[int]:
        style_name = (getattr(paragraph.style, "name", "") or "").strip().lower()
        if not style_name.startswith("heading"):
            return None

        match = re.search(r"(\d+)", style_name)
        if not match:
            return 1

        return max(1, int(match.group(1)))

    def table_text(table) -> str:
        rows = []
        for row in table.rows:
            cells = [
                " ".join((cell.text or "").split())
                for cell in row.cells
            ]
            row_text = " | ".join(cell for cell in cells if cell)
            if row_text:
                rows.append(row_text)
        return "\n".join(rows).strip()

    segments = []
    heading_stack = {}

    def current_heading_path() -> Optional[str]:
        path = [
            heading_stack[level]
            for level in sorted(heading_stack)
            if heading_stack.get(level)
        ]
        return " > ".join(path) if path else None

    for index, item in enumerate(iter_body_items(document), start=1):
        text = ""
        section = current_heading_path()
        is_heading = False

        if isinstance(item, Paragraph):
            text = paragraph_text(item)
            level = heading_level(item)
            if text and level is not None:
                heading_stack = {
                    existing_level: heading_text
                    for existing_level, heading_text in heading_stack.items()
                    if existing_level < level
                }
                heading_stack[level] = text
                section = current_heading_path()
                is_heading = True
        elif isinstance(item, Table):
            text = table_text(item)

        if not text:
            continue

        segments.append({
            "text": text,
            "section": section,
            "is_heading": is_heading,
        })

    if not any(segment["text"].strip() for segment in segments):
        raise DocumentExtractionError(
            "The uploaded DOCX does not contain extractable text."
        )

    blocks = []
    current_parts = []
    current_section_for_block = None
    current_size = 0
    target_block_chars = 1400

    def flush_block():
        nonlocal current_parts, current_section_for_block, current_size
        text = "\n\n".join(current_parts).strip()
        if text:
            blocks.append(
                ExtractedBlock(
                    text=text,
                    raw_text=text,
                    page=None,
                    section=current_section_for_block,
                    block_index=len(blocks) + 1,
                )
            )
        current_parts = []
        current_section_for_block = None
        current_size = 0

    for segment in segments:
        text = segment["text"]
        section = segment["section"]
        is_heading = segment["is_heading"]

        section_changed = (
            current_section_for_block is not None
            and section is not None
            and section != current_section_for_block
        )

        if (
            current_parts
            and (
                current_size + len(text) > target_block_chars
                or (is_heading and section_changed)
            )
        ):
            flush_block()

        if current_section_for_block is None:
            current_section_for_block = section

        current_parts.append(text)
        current_size += len(text)

    flush_block()

    if not blocks:
        joined_text = "\n\n".join(segment["text"] for segment in segments).strip()
        blocks.append(
            ExtractedBlock(
                text=joined_text,
                raw_text=joined_text,
                page=None,
                section=None,
                block_index=1,
            )
        )

    return ExtractedDocument(
        filename=filename,
        file_format="docx",
        file_size_bytes=len(file_bytes),
        blocks=blocks,
        pages_detected=None,
    )


def extract_pptx_document(file_bytes: bytes, filename: str) -> ExtractedDocument:
    try:
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE_TYPE
    except ImportError as exc:
        raise DocumentExtractionError(
            "PPTX support is not installed on this server."
        ) from exc

    try:
        presentation = Presentation(io.BytesIO(file_bytes))
    except Exception as exc:
        raise DocumentExtractionError(
            "The uploaded PPTX could not be opened. The file may be corrupted."
        ) from exc

    if len(presentation.slides) <= 0:
        raise DocumentExtractionError("The uploaded PPTX does not contain any slides.")

    def normalize_text(text: str) -> str:
        return " ".join((text or "").split()).strip()

    def text_frame_text(text_frame) -> str:
        lines = []
        for paragraph in text_frame.paragraphs:
            paragraph_text = normalize_text(paragraph.text)
            if not paragraph_text:
                continue

            level = max(0, int(getattr(paragraph, "level", 0) or 0))
            indent = "    " * level
            if level > 0:
                lines.append(f"{indent}• {paragraph_text}")
            else:
                lines.append(paragraph_text)

        return "\n".join(lines).strip()

    def table_text(table) -> str:
        rows = []
        for row in table.rows:
            cells = [
                normalize_text(cell.text)
                for cell in row.cells
            ]
            row_text = " | ".join(cell for cell in cells if cell)
            if row_text:
                rows.append(row_text)
        return "\n".join(rows).strip()

    def slide_title(slide, slide_number: int) -> str:
        title_shape = getattr(slide.shapes, "title", None)
        title = normalize_text(getattr(title_shape, "text", "") if title_shape else "")
        return title or f"Slide {slide_number}"

    def shape_text_blocks(slide, slide_title_shape) -> List[str]:
        blocks = []

        for shape in slide.shapes:
            if shape is slide_title_shape:
                continue

            if getattr(shape, "has_text_frame", False):
                text = text_frame_text(shape.text_frame)
                if text:
                    blocks.append(text)
                continue

            if getattr(shape, "has_table", False):
                text = table_text(shape.table)
                if text:
                    blocks.append(text)
                continue

            if getattr(shape, "shape_type", None) == MSO_SHAPE_TYPE.GROUP:
                for grouped_shape in shape.shapes:
                    if getattr(grouped_shape, "has_text_frame", False):
                        text = text_frame_text(grouped_shape.text_frame)
                        if text:
                            blocks.append(text)
                    elif getattr(grouped_shape, "has_table", False):
                        text = table_text(grouped_shape.table)
                        if text:
                            blocks.append(text)

        return blocks

    segments = []

    for slide_index, slide in enumerate(presentation.slides, start=1):
        title_shape = getattr(slide.shapes, "title", None)
        title = slide_title(slide, slide_index)
        slide_parts = [title]
        slide_parts.extend(shape_text_blocks(slide, title_shape))

        clean_slide_text = "\n\n".join(
            part for part in slide_parts if part and part.strip()
        ).strip()

        if not clean_slide_text:
            continue

        segments.append({
            "text": clean_slide_text,
            "raw_text": f"## SLIDE ##\n{title}\n\n{clean_slide_text}",
            "section": title,
            "is_new_slide": True,
        })

    if not any(segment["text"].strip() for segment in segments):
        raise DocumentExtractionError(
            "The uploaded PPTX does not contain extractable text."
        )

    blocks = []
    current_parts = []
    current_raw_parts = []
    current_section_for_block = None
    current_size = 0
    target_block_chars = 1400

    def flush_block():
        nonlocal current_parts, current_raw_parts, current_section_for_block, current_size
        text = "\n\n".join(current_parts).strip()
        raw_text = "\n\n".join(current_raw_parts).strip()
        if text:
            blocks.append(
                ExtractedBlock(
                    text=text,
                    raw_text=raw_text or text,
                    page=None,
                    section=current_section_for_block,
                    block_index=len(blocks) + 1,
                )
            )
        current_parts = []
        current_raw_parts = []
        current_section_for_block = None
        current_size = 0

    for segment in segments:
        text = segment["text"]
        raw_text = segment["raw_text"]
        section = segment["section"]

        if current_parts:
            flush_block()

        if len(text) <= target_block_chars:
            current_section_for_block = section
            current_parts.append(text)
            current_raw_parts.append(raw_text)
            current_size = len(text)
            flush_block()
            continue

        paragraphs = [part.strip() for part in text.split("\n\n") if part.strip()]
        raw_marker = f"## SLIDE ##\n{section}"
        current_section_for_block = section
        current_raw_parts.append(raw_marker)

        for paragraph in paragraphs:
            if current_parts and current_size + len(paragraph) > target_block_chars:
                flush_block()
                current_section_for_block = section
                current_raw_parts.append(raw_marker)

            current_parts.append(paragraph)
            current_raw_parts.append(paragraph)
            current_size += len(paragraph)

        flush_block()

    if not blocks:
        joined_text = "\n\n".join(segment["text"] for segment in segments).strip()
        joined_raw_text = "\n\n".join(segment["raw_text"] for segment in segments).strip()
        blocks.append(
            ExtractedBlock(
                text=joined_text,
                raw_text=joined_raw_text,
                page=None,
                section=None,
                block_index=1,
            )
        )

    return ExtractedDocument(
        filename=filename,
        file_format="pptx",
        file_size_bytes=len(file_bytes),
        blocks=blocks,
        pages_detected=None,
    )
