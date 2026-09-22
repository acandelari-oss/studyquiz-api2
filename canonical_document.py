"""Universal source-document contract for isolated structure mapping.

This module is not wired into production ingestion.  It preserves format-native
source fidelity at the edge and normalizes PDF, DOCX, and PPTX into one
ordered block representation for interpretation, anchoring, and segmentation.
"""

from __future__ import annotations

import io
import os
from dataclasses import dataclass, field, replace
from typing import Any, Dict, List, Optional

from document_extractors import DocumentExtractionError, PDF_PASSWORD_PROTECTED_MESSAGE


ROLE_TITLE = "title"
ROLE_HEADING = "heading"
ROLE_BODY = "body"
ROLE_LIST = "list"
ROLE_TABLE = "table"
ROLE_UNKNOWN = "unknown"


@dataclass(frozen=True, order=True)
class SourcePosition:
    unit_type: str
    unit_index: int
    block_index: int
    start_offset: int = 0
    end_offset: Optional[int] = None


@dataclass(frozen=True)
class CanonicalBlock:
    block_id: str
    text: str
    raw_text: str
    source_order: int
    source_position: SourcePosition
    role_hint: str = ROLE_UNKNOWN
    style_hint: Optional[str] = None
    native_hierarchy_hint: Optional[int] = None
    metadata: Dict[str, Any] = field(default_factory=dict)


@dataclass(frozen=True)
class CanonicalUnit:
    unit_id: str
    unit_type: str
    unit_index: int
    display_label: Optional[str]
    blocks: List[CanonicalBlock]
    metadata: Dict[str, Any] = field(default_factory=dict)


@dataclass(frozen=True)
class CanonicalDocumentInput:
    file_type: str
    filename: str
    title: str
    units: List[CanonicalUnit]
    metadata: Dict[str, Any] = field(default_factory=dict)

    @property
    def blocks(self) -> List[CanonicalBlock]:
        ordered = [block for unit in self.units for block in unit.blocks]
        return sorted(ordered, key=lambda block: block.source_order)


def adapt_document(file_bytes: bytes, filename: str) -> CanonicalDocumentInput:
    extension = os.path.splitext(filename or "")[1].lower().strip()
    if extension == ".pdf":
        return adapt_pdf_document(file_bytes, filename)
    if extension == ".docx":
        return adapt_docx_document(file_bytes, filename)
    if extension == ".pptx":
        return adapt_pptx_document(file_bytes, filename)
    raise DocumentExtractionError(
        "Unsupported document format. Please upload a PDF, DOCX, or PPTX file."
    )


def adapt_pdf_document(file_bytes: bytes, filename: str) -> CanonicalDocumentInput:
    from pypdf import PdfReader
    from pypdf.errors import DependencyError, FileNotDecryptedError, WrongPasswordError

    try:
        reader = PdfReader(io.BytesIO(file_bytes))
        if reader.is_encrypted and not reader.decrypt(""):
            raise DocumentExtractionError(PDF_PASSWORD_PROTECTED_MESSAGE)
    except DocumentExtractionError:
        raise
    except (FileNotDecryptedError, WrongPasswordError) as exc:
        raise DocumentExtractionError(PDF_PASSWORD_PROTECTED_MESSAGE) from exc
    except DependencyError as exc:
        raise DocumentExtractionError(
            "This PDF uses encryption that requires additional server support."
        ) from exc
    except Exception as exc:
        raise DocumentExtractionError(
            "The uploaded PDF could not be opened. The file may be corrupted."
        ) from exc

    units: List[CanonicalUnit] = []
    source_order = 0
    for page_index, page in enumerate(reader.pages, start=1):
        raw_text = page.extract_text() or ""
        page_width = _pdf_page_dimension(page, "width")
        page_height = _pdf_page_dimension(page, "height")
        page_blocks: List[CanonicalBlock] = []
        text_fragments = _pdf_text_fragments(page, page_width, page_height)
        for fragment_index, fragment in enumerate(text_fragments, start=1):
            fragment_text = fragment["text"]
            source_order += 1
            page_blocks.append(
                CanonicalBlock(
                    block_id=f"page-{page_index:04d}-fragment-{fragment_index:04d}",
                    text=" ".join(fragment_text.split()),
                    raw_text=fragment_text,
                    source_order=source_order,
                    source_position=SourcePosition(
                        unit_type="page",
                        unit_index=page_index,
                        block_index=fragment_index,
                        start_offset=0,
                        end_offset=len(fragment_text),
                    ),
                    role_hint=ROLE_BODY,
                    metadata={
                        **fragment["metadata"],
                        "page_index": page_index,
                        "page_width": page_width,
                        "page_height": page_height,
                        "pdf_text_unit": "text_fragment",
                    },
                )
            )
        units.append(
            CanonicalUnit(
                unit_id=f"page-{page_index:04d}",
                unit_type="page",
                unit_index=page_index,
                display_label=f"Page {page_index}",
                blocks=page_blocks,
                metadata={
                    "page_width": page_width,
                    "page_height": page_height,
                    "native_text_available": bool(text_fragments),
                },
            )
        )

    if not units:
        raise DocumentExtractionError("The uploaded PDF does not contain any pages.")

    return CanonicalDocumentInput(
        file_type="pdf",
        filename=filename,
        title=filename,
        units=units,
        metadata={"pages_detected": len(units)},
    )


def _pdf_page_dimension(page, dimension: str) -> Optional[float]:
    try:
        box = page.mediabox
        value = getattr(box, dimension)
        return float(value)
    except Exception:
        return None


def _pdf_text_fragments(
    page,
    page_width: Optional[float],
    page_height: Optional[float],
) -> List[Dict[str, Any]]:
    fragments: List[Dict[str, Any]] = []

    def visitor_text(text, current_transformation_matrix, text_matrix, font_dict, font_size):
        raw_text = text or ""
        if not raw_text.strip():
            return
        x, y = _pdf_text_position(current_transformation_matrix, text_matrix)
        metadata: Dict[str, Any] = {
            "x": x,
            "y": y,
            "x0": x,
            "y0": y,
            "font_size": float(font_size) if font_size is not None else None,
            "text_matrix": _pdf_matrix_values(text_matrix),
            "current_transformation_matrix": _pdf_matrix_values(current_transformation_matrix),
        }
        if page_width is not None:
            metadata["page_width"] = page_width
        if page_height is not None:
            metadata["page_height"] = page_height
        font_metadata = _pdf_font_metadata(font_dict)
        metadata.update(font_metadata)
        fragments.append({"text": raw_text, "metadata": metadata})

    try:
        page.extract_text(visitor_text=visitor_text)
    except Exception:
        return []
    return fragments


def _pdf_text_position(
    current_transformation_matrix,
    text_matrix,
) -> tuple[Optional[float], Optional[float]]:
    try:
        cm = [float(value) for value in current_transformation_matrix]
        tm = [float(value) for value in text_matrix]
        x = tm[4] * cm[0] + tm[5] * cm[2] + cm[4]
        y = tm[4] * cm[1] + tm[5] * cm[3] + cm[5]
        return x, y
    except Exception:
        try:
            tm = [float(value) for value in text_matrix]
            return tm[4], tm[5]
        except Exception:
            return None, None


def _pdf_matrix_values(matrix) -> List[float]:
    try:
        return [float(value) for value in matrix]
    except Exception:
        return []


def _pdf_font_metadata(font_dict) -> Dict[str, Any]:
    if not font_dict:
        return {}
    metadata: Dict[str, Any] = {}
    for source_key, target_key in (
        ("/BaseFont", "font_name"),
        ("/Subtype", "font_subtype"),
    ):
        try:
            value = font_dict.get(source_key)
            if value is not None:
                metadata[target_key] = str(value).lstrip("/")
        except Exception:
            pass
    return metadata


def adapt_docx_document(file_bytes: bytes, filename: str) -> CanonicalDocumentInput:
    try:
        from docx import Document
        from docx.table import Table
        from docx.text.paragraph import Paragraph
    except ImportError as exc:
        raise DocumentExtractionError(
            "DOCX support is not installed on this server."
        ) from exc

    try:
        document = Document(io.BytesIO(file_bytes))
    except Exception as exc:
        raise DocumentExtractionError(
            "The uploaded DOCX could not be opened. The file may be corrupted."
        ) from exc

    blocks: List[CanonicalBlock] = []
    source_order = 0
    numbering_index = _docx_numbering_index(document)

    for native_index, item in enumerate(_iter_docx_body_items(document), start=1):
        if isinstance(item, Paragraph):
            text = (item.text or "").strip()
            if not text:
                continue
            style_name = (getattr(item.style, "name", "") or "").strip() or None
            heading_level = _docx_heading_level(style_name)
            numbering = _docx_numbering_metadata(item, numbering_index)
            list_level = _docx_list_level(item, numbering)
            role_hint = ROLE_HEADING if heading_level is not None else ROLE_BODY
            if list_level is not None and heading_level is None:
                role_hint = ROLE_LIST
            hierarchy_hint = heading_level if heading_level is not None else list_level
            metadata: Dict[str, Any] = {"paragraph_index": native_index}
            if numbering:
                metadata["numbering"] = numbering
            source_order += 1
            blocks.append(
                CanonicalBlock(
                    block_id=f"paragraph-{native_index:04d}",
                    text=text,
                    raw_text=text,
                    source_order=source_order,
                    source_position=SourcePosition(
                        unit_type="document_flow",
                        unit_index=1,
                        block_index=native_index,
                        start_offset=0,
                        end_offset=len(text),
                    ),
                    role_hint=role_hint,
                    style_hint=style_name,
                    native_hierarchy_hint=hierarchy_hint,
                    metadata=metadata,
                )
            )
        elif isinstance(item, Table):
            text = _docx_table_text(item)
            if not text:
                continue
            source_order += 1
            blocks.append(
                CanonicalBlock(
                    block_id=f"table-{native_index:04d}",
                    text=text,
                    raw_text=text,
                    source_order=source_order,
                    source_position=SourcePosition(
                        unit_type="document_flow",
                        unit_index=1,
                        block_index=native_index,
                        start_offset=0,
                        end_offset=len(text),
                    ),
                    role_hint=ROLE_TABLE,
                    style_hint="table",
                    metadata={"body_item_index": native_index},
                )
            )

    if not blocks:
        raise DocumentExtractionError(
            "The uploaded DOCX does not contain extractable text."
        )

    return CanonicalDocumentInput(
        file_type="docx",
        filename=filename,
        title=filename,
        units=[
            CanonicalUnit(
                unit_id="document-flow-0001",
                unit_type="document_flow",
                unit_index=1,
                display_label="Document flow",
                blocks=blocks,
            )
        ],
    )


def adapt_pptx_document(file_bytes: bytes, filename: str) -> CanonicalDocumentInput:
    try:
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE_TYPE
    except ImportError as exc:
        raise DocumentExtractionError(
            "PPTX support is not installed on this server."
        ) from exc

    try:
        presentation = Presentation(io.BytesIO(file_bytes))
    except Exception as exc:
        raise DocumentExtractionError(
            "The uploaded PPTX could not be opened. The file may be corrupted."
        ) from exc

    if len(presentation.slides) <= 0:
        raise DocumentExtractionError("The uploaded PPTX does not contain any slides.")

    units: List[CanonicalUnit] = []
    source_order = 0
    for slide_index, slide in enumerate(presentation.slides, start=1):
        title_shape = getattr(slide.shapes, "title", None)
        layout_metadata = _pptx_slide_layout_metadata(presentation, slide)
        slide_blocks: List[CanonicalBlock] = []
        shape_counter = 0

        for shape in slide.shapes:
            shape_counter += 1
            for block_text, role_hint, hierarchy_hint, metadata in _pptx_shape_blocks(
                shape,
                title_shape,
            ):
                if not block_text:
                    continue
                source_order += 1
                block_index = len(slide_blocks) + 1
                slide_blocks.append(
                    CanonicalBlock(
                        block_id=f"slide-{slide_index:04d}-block-{block_index:04d}",
                        text=block_text,
                        raw_text=block_text,
                        source_order=source_order,
                        source_position=SourcePosition(
                            unit_type="slide",
                            unit_index=slide_index,
                            block_index=block_index,
                            start_offset=0,
                            end_offset=len(block_text),
                        ),
                        role_hint=role_hint,
                        style_hint=metadata.get("placeholder_type"),
                        native_hierarchy_hint=hierarchy_hint,
                        metadata={
                            **layout_metadata,
                            **metadata,
                            "shape_index": shape_counter,
                        },
                    )
                )

        if slide_blocks:
            title = next(
                (block.text for block in slide_blocks if block.role_hint == ROLE_TITLE),
                f"Slide {slide_index}",
            )
            units.append(
                CanonicalUnit(
                    unit_id=f"slide-{slide_index:04d}",
                    unit_type="slide",
                    unit_index=slide_index,
                    display_label=title,
                    blocks=slide_blocks,
                    metadata=layout_metadata,
                )
            )

    if not units:
        raise DocumentExtractionError(
            "The uploaded PPTX does not contain extractable text."
        )

    return CanonicalDocumentInput(
        file_type="pptx",
        filename=filename,
        title=filename,
        units=units,
        metadata={"slides_detected": len(presentation.slides)},
    )


def canonical_document_from_blocks(
    *,
    file_type: str,
    filename: str,
    units: List[CanonicalUnit],
) -> CanonicalDocumentInput:
    return CanonicalDocumentInput(
        file_type=file_type,
        filename=filename,
        title=filename,
        units=units,
    )


def block_with_text(block: CanonicalBlock, text: str) -> CanonicalBlock:
    return replace(
        block,
        text=text,
        raw_text=text,
        source_position=replace(
            block.source_position,
            start_offset=0,
            end_offset=len(text),
        ),
    )


def _iter_docx_body_items(document):
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    for child in document.element.body.iterchildren():
        if child.tag.endswith("}p"):
            yield Paragraph(child, document)
        elif child.tag.endswith("}tbl"):
            yield Table(child, document)


def _docx_heading_level(style_name: Optional[str]) -> Optional[int]:
    if not style_name:
        return None
    lowered = style_name.lower()
    if not lowered.startswith("heading"):
        return None
    digits = "".join(char for char in lowered if char.isdigit())
    return int(digits) if digits else 1


def _docx_list_level(paragraph, numbering: Optional[Dict[str, Any]] = None) -> Optional[int]:
    if numbering and "ilvl" in numbering:
        return numbering["ilvl"]
    try:
        num_pr = paragraph._p.pPr.numPr  # noqa: SLF001 - native hint extraction.
        if num_pr is None or num_pr.ilvl is None:
            return None
        return int(num_pr.ilvl.val)
    except Exception:
        return None


def _docx_numbering_metadata(
    paragraph,
    numbering_index: Optional[Dict[str, Any]] = None,
) -> Optional[Dict[str, Any]]:
    """Preserve Word-native numbering evidence without rendering labels."""

    numbering_index = numbering_index or {}
    direct = _docx_num_pr_metadata(_paragraph_num_pr(paragraph))
    if direct:
        direct["source"] = "direct"
        _docx_apply_numbering_definition(direct, numbering_index)
        return direct

    style = getattr(paragraph, "style", None)
    style_numbering = _docx_num_pr_metadata(_style_num_pr(style))
    if style_numbering:
        style_numbering["source"] = "style"
        style_id = getattr(style, "style_id", None)
        if style_id:
            style_numbering["style_id"] = style_id
        _docx_apply_numbering_definition(style_numbering, numbering_index)
        return style_numbering

    style_id = getattr(style, "style_id", None)
    if style_id:
        linked = numbering_index.get("style_links", {}).get(style_id)
        if linked:
            return {
                **linked,
                "source": "style_linked",
                "style_id": style_id,
            }
    return None


def _paragraph_num_pr(paragraph):
    try:
        p_pr = paragraph._p.pPr  # noqa: SLF001 - native evidence extraction.
        return p_pr.numPr if p_pr is not None else None
    except Exception:
        return None


def _style_num_pr(style):
    try:
        element = getattr(style, "element", None)
        p_pr = element.pPr if element is not None else None
        return p_pr.numPr if p_pr is not None else None
    except Exception:
        return None


def _docx_num_pr_metadata(num_pr) -> Optional[Dict[str, Any]]:
    if num_pr is None:
        return None
    metadata: Dict[str, Any] = {}
    try:
        if num_pr.numId is not None:
            metadata["num_id"] = _docx_scalar(num_pr.numId.val)
    except Exception:
        pass
    try:
        if num_pr.ilvl is not None:
            metadata["ilvl"] = int(num_pr.ilvl.val)
    except Exception:
        pass
    return metadata or None


def _docx_numbering_index(document) -> Dict[str, Any]:
    try:
        numbering = document.part.numbering_part.element
    except Exception:
        return {"nums": {}, "abstracts": {}, "style_links": {}}

    nums: Dict[Any, Dict[str, Any]] = {}
    abstracts: Dict[Any, Dict[str, Any]] = {}
    style_links: Dict[str, Dict[str, Any]] = {}

    for abstract in numbering.findall(_docx_qn("w:abstractNum")):
        abstract_num_id = _docx_scalar(abstract.get(_docx_qn("w:abstractNumId")))
        levels: Dict[int, Dict[str, Any]] = {}
        for level in abstract.findall(_docx_qn("w:lvl")):
            ilvl_value = _docx_scalar(level.get(_docx_qn("w:ilvl")))
            if not isinstance(ilvl_value, int):
                continue
            level_metadata: Dict[str, Any] = {
                "abstract_num_id": abstract_num_id,
                "ilvl": ilvl_value,
            }
            _docx_set_child_value(level_metadata, "format", level, "w:numFmt")
            _docx_set_child_value(level_metadata, "level_text", level, "w:lvlText")
            _docx_set_child_value(level_metadata, "start", level, "w:start")
            style_id = _docx_child_value(level, "w:pStyle")
            if style_id:
                level_metadata["style_id"] = style_id
                style_links[str(style_id)] = dict(level_metadata)
            levels[ilvl_value] = level_metadata
        abstracts[abstract_num_id] = {"levels": levels}

    for num in numbering.findall(_docx_qn("w:num")):
        num_id = _docx_scalar(num.get(_docx_qn("w:numId")))
        abstract_num_id = _docx_child_value(num, "w:abstractNumId")
        start_overrides: Dict[int, Any] = {}
        for override in num.findall(_docx_qn("w:lvlOverride")):
            ilvl_value = _docx_scalar(override.get(_docx_qn("w:ilvl")))
            if isinstance(ilvl_value, int):
                start_override = _docx_child_value(override, "w:startOverride")
                if start_override is not None:
                    start_overrides[ilvl_value] = start_override
        nums[num_id] = {
            "abstract_num_id": abstract_num_id,
            "start_overrides": start_overrides,
        }

    return {"nums": nums, "abstracts": abstracts, "style_links": style_links}


def _docx_apply_numbering_definition(
    metadata: Dict[str, Any],
    numbering_index: Dict[str, Any],
) -> None:
    num_id = metadata.get("num_id")
    num_definition = numbering_index.get("nums", {}).get(num_id)
    if num_definition:
        metadata.setdefault("abstract_num_id", num_definition.get("abstract_num_id"))
    abstract_num_id = metadata.get("abstract_num_id")
    ilvl = metadata.get("ilvl")
    level_definition = (
        numbering_index.get("abstracts", {})
        .get(abstract_num_id, {})
        .get("levels", {})
        .get(ilvl)
    )
    if level_definition:
        for key in ("format", "level_text", "start", "style_id"):
            if key in level_definition:
                metadata.setdefault(key, level_definition[key])
    if num_definition and isinstance(ilvl, int):
        start_override = num_definition.get("start_overrides", {}).get(ilvl)
        if start_override is not None:
            metadata["start_override"] = start_override


def _docx_set_child_value(
    target: Dict[str, Any],
    target_key: str,
    parent,
    child_tag: str,
) -> None:
    value = _docx_child_value(parent, child_tag)
    if value is not None:
        target[target_key] = value


def _docx_child_value(parent, child_tag: str) -> Optional[Any]:
    child = parent.find(_docx_qn(child_tag))
    if child is None:
        return None
    return _docx_scalar(child.get(_docx_qn("w:val")))


def _docx_qn(tag: str) -> str:
    from docx.oxml.ns import qn

    return qn(tag)


def _docx_scalar(value: Any) -> Any:
    if value is None:
        return None
    text = str(value)
    if text.isdigit():
        return int(text)
    return text


def _docx_table_text(table) -> str:
    rows = []
    for row in table.rows:
        cells = [" ".join((cell.text or "").split()) for cell in row.cells]
        row_text = " | ".join(cell for cell in cells if cell)
        if row_text:
            rows.append(row_text)
    return "\n".join(rows).strip()


def _pptx_slide_layout_metadata(presentation, slide) -> Dict[str, Any]:
    layout = getattr(slide, "slide_layout", None)
    if layout is None:
        return {}
    metadata: Dict[str, Any] = {}
    layout_name = (getattr(layout, "name", "") or "").strip()
    if layout_name:
        metadata["layout_name"] = layout_name
    try:
        for index, candidate in enumerate(presentation.slide_layouts):
            if getattr(candidate, "element", None) is getattr(layout, "element", None):
                metadata["layout_index"] = index
                break
    except Exception:
        pass
    return metadata


def _pptx_shape_blocks(shape, title_shape) -> List[tuple[str, str, Optional[int], Dict[str, Any]]]:
    try:
        from pptx.enum.shapes import MSO_SHAPE_TYPE
    except ImportError:  # pragma: no cover - import handled by caller.
        return []

    if getattr(shape, "shape_type", None) == MSO_SHAPE_TYPE.GROUP:
        blocks = []
        for grouped_shape in shape.shapes:
            blocks.extend(_pptx_shape_blocks(grouped_shape, title_shape))
        return blocks

    placeholder_type = None
    placeholder_idx = None
    if getattr(shape, "is_placeholder", False):
        try:
            placeholder_type = str(shape.placeholder_format.type)
        except Exception:
            placeholder_type = "placeholder"
        try:
            placeholder_idx = int(shape.placeholder_format.idx)
        except Exception:
            placeholder_idx = None

    shape_metadata = _pptx_shape_metadata(shape)
    if placeholder_type is not None:
        shape_metadata["placeholder_type"] = placeholder_type
    if placeholder_idx is not None:
        shape_metadata["placeholder_idx"] = placeholder_idx

    if getattr(shape, "has_table", False):
        text = _pptx_table_text(shape.table)
        return [(text, ROLE_TABLE, None, shape_metadata)] if text else []

    if not getattr(shape, "has_text_frame", False):
        return []

    role_hint = ROLE_TITLE if _same_pptx_shape(shape, title_shape) else ROLE_BODY
    paragraphs = []
    for paragraph in shape.text_frame.paragraphs:
        paragraph_text = " ".join((paragraph.text or "").split()).strip()
        if not paragraph_text:
            continue
        level = max(0, int(getattr(paragraph, "level", 0) or 0))
        paragraph_role = role_hint if role_hint == ROLE_TITLE else (ROLE_LIST if level > 0 else ROLE_BODY)
        paragraphs.append(
            (
                paragraph_text,
                paragraph_role,
                level,
                {**shape_metadata, "bullet_level": level},
            )
        )
    return paragraphs


def _same_pptx_shape(shape, other_shape) -> bool:
    if other_shape is None:
        return False
    try:
        return getattr(shape, "element", None) is getattr(other_shape, "element", None)
    except Exception:
        return False


def _pptx_shape_metadata(shape) -> Dict[str, Any]:
    metadata: Dict[str, Any] = {}
    name = (getattr(shape, "name", "") or "").strip()
    if name:
        metadata["shape_name"] = name
    for key, attr in (
        ("x", "left"),
        ("y", "top"),
        ("width", "width"),
        ("height", "height"),
    ):
        try:
            metadata[key] = int(getattr(shape, attr))
        except Exception:
            pass
    return metadata


def _pptx_table_text(table) -> str:
    rows = []
    for row in table.rows:
        cells = [" ".join((cell.text or "").split()) for cell in row.cells]
        row_text = " | ".join(cell for cell in cells if cell)
        if row_text:
            rows.append(row_text)
    return "\n".join(rows).strip()
